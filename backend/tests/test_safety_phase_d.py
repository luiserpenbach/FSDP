"""Safety phase D: safety review packages (PDF + XLSX, change log since the
previous package), the safety approver grant on acceptance and release, and
the certification evidence view with its gaps."""

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.config import settings as app_settings
from app.main import app


@pytest.fixture(autouse=True)
def _package_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_settings, "safety_files_dir", str(tmp_path / "packages"))


def _project(client: TestClient) -> str:
    return client.post("/projects", json={"name": "LOX GSE", "part_name_prefix": "LOX"}).json()[
        "id"
    ]


def _item(item_id: str, category: str, symbol_key: str, tag: str, **fields: object) -> dict:
    return {
        "item_id": item_id,
        "kind": "symbol",
        "category": category,
        "symbol_key": symbol_key,
        "symbol_name": symbol_key.replace("_", " ").title(),
        "tag": tag,
        "label": None,
        "zone": "C-4",
        "x": 10,
        "y": 10,
        "part_id": None,
        "dnp": False,
        "spare": 0,
        "fields": {"service": "LOX", "line_number": "L-2014", **fields},
    }


def _drawing(client: TestClient, project_id: str) -> tuple[str, str]:
    drawing = client.post(
        f"/projects/{project_id}/drawings", json={"title": "LOX P&ID", "size": "ANSI_E"}
    ).json()
    sheet_id = drawing["sheets"][0]["id"]
    items = [
        _item("fv201", "valve", "pneumatic_valve", "FV-201"),
        _item("pt205", "instrument", "instrument", "PT-205"),
    ]
    saved = client.put(f"/sheets/{sheet_id}", json={"index": {"items": items, "lines": []}})
    assert saved.status_code == 200, saved.text
    client.post(
        f"/drawings/{drawing['id']}/revisions", json={"label": "A", "description": "Issued"}
    )
    return drawing["id"], sheet_id


def _hazard(client: TestClient, project_id: str, **extra: object) -> dict:
    body = {
        "title": "Fast-fill overpressure",
        "category": "overpressure",
        "severity_initial": "I",
        "likelihood_initial": "C",
    }
    body.update(extra)
    response = client.post(f"/projects/{project_id}/hazards", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def _requirement(client: TestClient, project_id: str, key: str, **extra: object) -> dict:
    body = {
        "project_id": project_id,
        "key": key,
        "title": f"Requirement {key}",
        "text": "Shall.",
        "requirement_type": "safety",
        "category": "safety",
        "safety_critical": True,
    }
    body.update(extra)
    response = client.post("/requirements", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def _released_worksheet(
    client: TestClient, project_id: str, drawing_id: str, hazard_id: str
) -> dict:
    worksheet = client.post(
        f"/projects/{project_id}/fmea", json={"title": "LOX fill", "drawing_id": drawing_id}
    ).json()
    client.post(f"/fmea/{worksheet['id']}/generate", json={"categories": ["valve"]})
    rows = client.get(f"/fmea/{worksheet['id']}/rows").json()
    patches = []
    for row in rows:
        patch = {"id": row["id"], "occurrence": 3, "detection": 2, "hazard_id": hazard_id}
        if row["detection_kind"] == "none":
            patch["detection_reason"] = "covered by relief"
        patches.append(patch)
    bulk = client.post(f"/fmea/{worksheet['id']}/rows/bulk", json={"rows": patches})
    assert bulk.status_code == 200, bulk.text
    released = client.post(f"/fmea/{worksheet['id']}/release", json={"note": "Design review 1"})
    assert released.status_code == 201, released.text
    return worksheet


def _engineer(client: TestClient, email: str) -> TestClient:
    created = client.post(
        "/auth/users",
        json={
            "email": email,
            "name": email.split("@")[0],
            "password": "eng-password-1",
            "role": "engineer",
        },
    )
    assert created.status_code == 201, created.text
    other = TestClient(app)
    login = other.post("/auth/login", json={"email": email, "password": "eng-password-1"})
    assert login.status_code == 200, login.text
    return other


def test_package_generates_pdf_xlsx_and_change_log(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, _sheet_id = _drawing(client, project_id)
    hazard = _hazard(client, project_id)
    requirement = _requirement(client, project_id, "REQ-SAF-001")
    linked = client.post(
        f"/hazards/{hazard['id']}/controls", json={"type": "requirement", "id": requirement["id"]}
    )
    assert linked.status_code == 201, linked.text
    worksheet = _released_worksheet(client, project_id, drawing_id, hazard["id"])

    created = client.post(
        f"/projects/{project_id}/safety/packages", json={"title": "PDR safety package"}
    )
    assert created.status_code == 201, created.text
    package = created.json()
    assert package["title"] == "PDR safety package"
    assert package["change_log_from"] is None
    assert package["scope"]["drawings"][0]["revision"] == "A"
    assert package["scope"]["worksheets"][0]["id"] == worksheet["id"]
    summary = package["summary"]
    assert summary["hazards"] == 1
    assert summary["hazards_high_open"] == 1
    assert summary["requirements_safety"] == 1
    assert summary["requirements_verified"] == 0
    assert summary["fmea_rows"] > 0
    assert summary["actions"] >= 2  # the open hazard and the unverified requirement

    pdf = client.get(package["pdf_url"])
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    reader = PdfReader(__import__("io").BytesIO(pdf.content))
    assert len(reader.pages) >= 9  # cover, two matrices, hazards, FMEA, matrix, actions, DRC, log

    xlsx = client.get(package["xlsx_url"])
    assert xlsx.status_code == 200
    workbook = load_workbook(__import__("io").BytesIO(xlsx.content))
    assert workbook.sheetnames == [
        "Summary",
        "Hazards",
        "FMEA",
        "Requirements",
        "Open actions",
        "Design rules",
        "Change log",
    ]
    hazard_rows = list(workbook["Hazards"].iter_rows(values_only=True))
    assert hazard_rows[1][0] == hazard["key"]
    assert workbook["FMEA"].max_row == summary["fmea_rows"] + 1

    # A change after the first package lands in the second package's change log only.
    client.put(f"/hazards/{hazard['id']}", json={"owner": "K. Ortega"})
    second = client.post(f"/projects/{project_id}/safety/packages", json={}).json()
    assert second["title"].startswith("Safety review package 2")
    assert second["change_log_from"] == package["created_at"]
    assert second["summary"]["changes"] >= 1
    listed = client.get(f"/projects/{project_id}/safety/packages").json()
    assert [entry["id"] for entry in listed] == [second["id"], package["id"]]

    deleted = client.delete(f"/safety/packages/{package['id']}")
    assert deleted.status_code == 204
    assert client.get(package["pdf_url"]).status_code == 404
    assert client.get("/safety/packages/nope/pdf").status_code == 404
    assert client.get(f"/safety/packages/{second['id']}/docx").status_code == 400


def test_approver_grant_gates_accept_and_release(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, _sheet_id = _drawing(client, project_id)
    hazard = _hazard(client, project_id)
    engineer = _engineer(client, "eng@fsdp.test")
    approver = _engineer(client, "approver@fsdp.test")

    # No approvers configured: any engineer may accept and release.
    open_release = engineer.post(f"/hazards/{hazard['id']}/accept", json={"justification": "ok"})
    assert open_release.status_code == 200, open_release.text

    settings = client.get(f"/projects/{project_id}/safety-settings").json()["settings"]
    settings["approvers"] = ["approver@fsdp.test"]
    saved = client.put(f"/projects/{project_id}/safety-settings", json={"settings": settings})
    assert saved.status_code == 200, saved.text

    second = _hazard(client, project_id, title="Second hazard")
    forbidden = engineer.post(f"/hazards/{second['id']}/accept", json={"justification": "ok"})
    assert forbidden.status_code == 403
    allowed = approver.post(f"/hazards/{second['id']}/accept", json={"justification": "ok"})
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["accepted_by"] == "approver@fsdp.test"

    worksheet = client.post(
        f"/projects/{project_id}/fmea", json={"title": "LOX fill", "drawing_id": drawing_id}
    ).json()
    client.post(f"/fmea/{worksheet['id']}/generate", json={"categories": ["valve"]})
    rows = client.get(f"/fmea/{worksheet['id']}/rows").json()
    patches = [
        {
            "id": row["id"],
            "occurrence": 2,
            "detection": 2,
            "hazard_id": hazard["id"],
            "detection_reason": "relief",
        }
        for row in rows
    ]
    client.post(f"/fmea/{worksheet['id']}/rows/bulk", json={"rows": patches})
    blocked = engineer.post(f"/fmea/{worksheet['id']}/release", json={})
    assert blocked.status_code == 403
    admin_ok = client.post(f"/fmea/{worksheet['id']}/release", json={})
    assert admin_ok.status_code == 201, admin_ok.text


def test_certification_evidence_lists_evidence_and_gaps(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, sheet_id = _drawing(client, project_id)
    hazard = _hazard(client, project_id)
    verified = _requirement(client, project_id, "REQ-SAF-001")
    pending = _requirement(client, project_id, "REQ-SAF-002")
    for requirement in (verified, pending):
        client.post(
            f"/hazards/{hazard['id']}/controls",
            json={"type": "requirement", "id": requirement["id"]},
        )
    client.post(
        f"/requirements/{verified['id']}/evidence",
        json={"kind": "test", "ref_type": "report", "ref_id": "TR-1", "status": "pass"},
    )
    worksheet = _released_worksheet(client, project_id, drawing_id, hazard["id"])

    before = client.get(f"/projects/{project_id}/certification/evidence").json()
    assert before["ready"] is False
    kinds = {gap["kind"] for gap in before["gaps"]}
    assert kinds == {"hazard_not_accepted", "requirement_not_verified", "no_package"}
    assert [entry["key"] for entry in before["verified_requirements"]] == [verified["key"]]
    assert before["released_worksheets"][0]["id"] == worksheet["id"]
    assert before["released_worksheets"][0]["behind_drawing"] is False
    assert any(gap["key"] == pending["key"] for gap in before["gaps"])

    # Accept the hazard, verify the requirement, generate a package: gaps close.
    client.post(f"/hazards/{hazard['id']}/accept", json={"justification": "Residual risk agreed."})
    client.post(
        f"/requirements/{pending['id']}/evidence",
        json={"kind": "waiver", "ref_type": "waiver", "ref_id": "WV-7", "status": "pass"},
    )
    client.post(f"/projects/{project_id}/safety/packages", json={})
    after = client.get(f"/projects/{project_id}/certification/evidence").json()
    assert after["ready"] is True, after["gaps"]
    assert after["counts"]["packages"] == 1
    assert [entry["key"] for entry in after["accepted_hazards"]] == [hazard["key"]]

    # A new drawing revision leaves the release behind the drawing.
    client.post(f"/drawings/{drawing_id}/revisions", json={"label": "B", "description": "Rev B"})
    behind = client.get(f"/projects/{project_id}/certification/evidence").json()
    assert behind["ready"] is False
    assert {gap["kind"] for gap in behind["gaps"]} == {"worksheet_behind_drawing"}
    assert behind["released_worksheets"][0]["current_drawing_revision"] == "B"
    assert sheet_id  # the sheet is unchanged; only the revision label moved
