"""Bulk import of parts and requirements from CSV, XLSX and pasted rows (Excel parity, §2.4)."""

import csv
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.services.bulk_import import (
    MAX_IMPORT_BYTES,
    PART_SPEC,
    REQUIREMENT_SPEC,
    resolve_mapping,
    table_from_upload,
)

PART_HEADER = "Part Number,Description,Type,Mfr,Material,Pressure Rating (bar),Lifecycle"


def _csv(*lines: str) -> bytes:
    return ("\n".join(lines) + "\n").encode()


def _upload(
    client: TestClient,
    path: str,
    content: bytes,
    filename: str = "parts.csv",
    **params,
):
    return client.post(path, params=params, files={"file": (filename, content, "text/csv")})


def _import_parts(client: TestClient, content: bytes, filename: str = "parts.csv", **params):
    return _upload(client, "/parts/import", content, filename, **params)


def _part(client: TestClient, number: str = "PV-1001", **extra) -> dict:
    body = {
        "part_number": number,
        "description": f"Valve {number}",
        "part_type": "valve",
        "material": "316L",
        "pressure_rating_bar": 200,
        "manufacturer": "Acme",
        **extra,
    }
    created = client.post("/parts", json=body)
    assert created.status_code == 201, created.text
    return created.json()


def _project(client: TestClient, name: str = "AMB2") -> str:
    return client.post("/projects", json={"name": name}).json()["id"]


def _viewer(client: TestClient) -> TestClient:
    created = client.post(
        "/auth/users",
        json={
            "email": "viewer@fsdp.test",
            "name": "Viewer",
            "password": "viewer-password",
            "role": "viewer",
        },
    )
    assert created.status_code == 201, created.text
    viewer = TestClient(app)
    login = viewer.post(
        "/auth/login", json={"email": "viewer@fsdp.test", "password": "viewer-password"}
    )
    assert login.status_code == 200
    return viewer


def test_header_aliases_map_case_space_and_underscore_insensitively() -> None:
    mapping = resolve_mapping(
        PART_SPEC,
        ["PN", "Mfr", " pressure_rating (BAR) ", "Lifecycle Status", "Qual", "Colour", "Part #"],
    )
    assert mapping.mapping == {
        "PN": "part_number",
        "Mfr": "manufacturer",
        " pressure_rating (BAR) ": "pressure_rating_bar",
        "Lifecycle Status": "lifecycle_status",
        "Qual": "qualification_status",
        "Colour": None,
        "Part #": None,
    }
    assert mapping.warnings == ["Column 'Part #' also maps to part_number; using column 'PN'"]

    for header in ("Part Number", "part_no", "PART-NUMBER", "p/n"):
        assert resolve_mapping(PART_SPEC, [header]).mapping[header] == "part_number"
    for header, expected in (
        ("Manufacturer", "manufacturer"),
        ("Pressure Rating", "pressure_rating_bar"),
        ("Name", "part_number"),  # the catalog CSV export's header
        ("Completeness", None),  # computed, never imported
    ):
        assert resolve_mapping(PART_SPEC, [header]).mapping[header] == expected

    requirement = resolve_mapping(
        REQUIREMENT_SPEC, ["ID", "Title", "Description", "Verification Method", "Type", "Owner"]
    )
    assert requirement.mapping == {
        "ID": "key",
        "Title": "title",
        "Description": "text",
        "Verification Method": "verification_method",
        "Type": "requirement_type",
        "Owner": "owner",
    }


def test_csv_dry_run_reports_plan_and_writes_nothing(client: TestClient) -> None:
    content = _csv(
        PART_HEADER + ",Colour",
        "PV-1001,Ball valve,valve,Acme,316L,206,active,red",
        "",
        "RG-2001,Regulator,regulator,,Brass,40.5,,blue",
    )
    response = _import_parts(client, content)
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["dry_run"] is True and report["committed"] is False
    assert report["mapping"] == {
        "Part Number": "part_number",
        "Description": "description",
        "Type": "part_type",
        "Mfr": "manufacturer",
        "Material": "material",
        "Pressure Rating (bar)": "pressure_rating_bar",
        "Lifecycle": "lifecycle_status",
        "Colour": None,
    }
    assert report["summary"] == {"create": 2, "update": 0, "unchanged": 0, "error": 0}
    # Data rows are numbered from 1 below the header; the blank row keeps its number.
    assert [(row["row"], row["action"], row["key"]) for row in report["rows"]] == [
        (1, "create", "PV-1001"),
        (3, "create", "RG-2001"),
    ]
    first = report["rows"][0]["changes"]
    assert first["pressure_rating_bar"] == [None, 206.0]
    assert first["lifecycle_status"] == [None, "active"]
    assert report["rows"][1]["changes"]["lifecycle_status"] == [None, "draft"]
    assert client.get("/parts").json() == []


def test_commit_creates_parts_and_records_changes(client: TestClient) -> None:
    content = _csv(PART_HEADER, "PV-1001,Ball valve,valve,Acme,316L,206,active")
    response = _import_parts(client, content, dry_run="false")
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["committed"] is True
    created_id = report["rows"][0]["id"]
    parts = client.get("/parts").json()
    assert [part["id"] for part in parts] == [created_id]
    assert parts[0]["pressure_rating_bar"] == 206
    assert parts[0]["manufacturer"] == "Acme"

    changes = client.get("/changes").json()
    by_object = {(change["object_type"], change["object_id"]): change for change in changes}
    assert by_object[("part", created_id)]["action"] == "created"
    assert by_object[("part", created_id)]["actor"] == "engineer@fsdp.test"
    summary = by_object[("catalog", "parts")]
    assert summary["action"] == "imported"
    assert "parts.csv" in summary["summary"] and "1 created" in summary["summary"]
    # The new part type is remembered in the catalog settings like a single create.
    assert "valve" in client.get("/catalog/settings").json()["part_types"]


def test_bom_prefixed_semicolon_csv_from_european_excel(client: TestClient) -> None:
    text = "Part Number;Description;Type;Pressure Rating (bar);Mass (kg);Preferred\r\n"
    text += "PV-1001;Ventil, Kugelhahn;valve;2,5;1.234,5;x\r\n"
    content = text.encode("utf-8-sig")
    assert content.startswith(b"\xef\xbb\xbf")
    report = _import_parts(client, content, filename="teile.csv").json()
    assert report["mapping"]["Part Number"] == "part_number"
    row = report["rows"][0]
    assert row["action"] == "create", row
    assert row["changes"]["description"] == [None, "Ventil, Kugelhahn"]
    assert row["changes"]["pressure_rating_bar"] == [None, 2.5]
    assert row["changes"]["mass_kg"] == [None, 1234.5]
    assert row["changes"]["preferred"] == [None, True]


def test_delimiter_detection_falls_back_and_cp1252_decodes() -> None:
    content = "Part Number\tDescription\nPV-1\tVálvula\n".encode("cp1252")
    table = table_from_upload(content, "parts.txt", PART_SPEC.key_aliases)
    assert table.headers == ["Part Number", "Description"]
    assert table.rows == [(1, ["PV-1", "Válvula"])]
    single = table_from_upload(b"part_number\nPV-1\n", "parts.csv", PART_SPEC.key_aliases)
    assert single.headers == ["part_number"]


def test_xlsx_upload_with_title_block_and_native_cell_types(client: TestClient) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Parts catalog export"])
    sheet.append([])
    sheet.append(["Part No", "Description", "Type", "Rev", "Bar", "Temp Min", "Preferred"])
    sheet.append(["PV-1001", "Ball valve", "valve", 2, 206.0, -40, True])
    sheet.append(["PV-1002", "Needle valve", "valve", "B", 413.7, None, False])
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = _import_parts(client, buffer.getvalue(), filename="parts.xlsx", dry_run="false")
    assert response.status_code == 200, response.text
    assert response.json()["summary"]["create"] == 2
    parts = {part["part_number"]: part for part in client.get("/parts").json()}
    assert parts["PV-1001"]["revision"] == "2"
    assert parts["PV-1001"]["pressure_rating_bar"] == 206
    assert parts["PV-1001"]["temperature_min_c"] == -40
    assert parts["PV-1001"]["preferred"] is True
    assert parts["PV-1002"]["pressure_rating_bar"] == 413.7
    assert parts["PV-1002"]["preferred"] is False


def test_unreadable_files_are_rejected(client: TestClient) -> None:
    bad_xlsx = _import_parts(client, b"not a zip", filename="parts.xlsx")
    assert bad_xlsx.status_code == 422
    assert "xlsx" in bad_xlsx.json()["detail"]
    legacy = _import_parts(client, b"\xd0\xcf\x11\xe0", filename="parts.xls")
    assert legacy.status_code == 422
    no_key = _import_parts(client, _csv("Description,Type", "Valve,valve"))
    assert no_key.status_code == 422
    assert "part_number" in no_key.json()["detail"]
    empty = _import_parts(client, b"")
    assert empty.status_code == 422


def test_upsert_diffs_existing_rows_and_detects_unchanged(client: TestClient) -> None:
    changed = _part(client, "PV-1001", material="316L", pressure_rating_bar=200)
    same = _part(client, "PV-1002", material="Brass")
    content = _csv(
        "part_number,material,pressure_rating_bar,notes",
        "pv-1001 ,316,250,",  # case/space-insensitive match; blank notes = leave unchanged
        "PV-1002,Brass,200,",
        "PV-1003,Monel,100,",
    )
    response = _import_parts(client, content, mode="upsert")
    assert response.status_code == 200, response.text
    rows = response.json()["rows"]
    assert rows[0]["action"] == "update"
    assert rows[0]["id"] == changed["id"]
    assert rows[0]["changes"] == {
        "material": ["316L", "316"],
        "pressure_rating_bar": [200.0, 250.0],
    }
    assert rows[1] == {
        "row": 2,
        "action": "unchanged",
        "key": "PV-1002",
        "id": same["id"],
        "errors": [],
        "changes": {},
    }
    assert rows[2]["action"] == "error"
    assert rows[2]["errors"] == [
        {"field": "description", "message": "is required"},
        {"field": "part_type", "message": "is required"},
    ]

    content = _csv(
        "part_number,material,pressure_rating_bar",
        "pv-1001,316,250",
        "PV-1002,Brass,200",
    )
    committed = _import_parts(client, content, mode="upsert", dry_run="false")
    assert committed.status_code == 200, committed.text
    assert committed.json()["summary"] == {"create": 0, "update": 1, "unchanged": 1, "error": 0}
    updated = client.get(f"/parts/{changed['id']}").json()
    assert updated["material"] == "316"
    assert updated["pressure_rating_bar"] == 250
    assert updated["part_number"] == "PV-1001"  # never renamed by a case-only difference


def test_create_only_refuses_existing_keys(client: TestClient) -> None:
    _part(client, "PV-1001")
    content = _csv("part_number,description,part_type", "PV-1001,Valve,valve")
    row = _import_parts(client, content).json()["rows"][0]
    assert row["action"] == "error"
    assert row["errors"][0]["field"] == "part_number"
    assert "mode=upsert" in row["errors"][0]["message"]


def test_commit_is_all_or_nothing(client: TestClient) -> None:
    _part(client, "PV-1001", material="316L")
    content = _csv(
        "part_number,description,part_type,material,pressure_rating_bar",
        "PV-1001,Valve,valve,Monel,",
        "PV-2001,Good,valve,316L,100",
        "PV-2002,Bad,valve,316L,lots",
    )
    response = _import_parts(client, content, mode="upsert", dry_run="false")
    assert response.status_code == 422
    report = response.json()
    assert report["committed"] is False
    assert report["summary"] == {"create": 1, "update": 1, "unchanged": 0, "error": 1}
    bad = report["rows"][2]
    assert bad["errors"][0]["field"] == "pressure_rating_bar"
    assert "(got 'lots')" in bad["errors"][0]["message"]
    parts = client.get("/parts").json()
    assert [(part["part_number"], part["material"]) for part in parts] == [("PV-1001", "316L")]
    assert not [c for c in client.get("/changes").json() if c["object_type"] == "catalog"]


def test_duplicate_keys_within_the_file_are_errors(client: TestClient) -> None:
    content = _csv(
        "part_number,description,part_type",
        "PV-1001,Valve,valve",
        "pv-1001,Same valve again,valve",
    )
    rows = _import_parts(client, content).json()["rows"]
    assert rows[0]["action"] == "create"
    assert rows[1]["action"] == "error"
    assert rows[1]["errors"] == [
        {"field": "part_number", "message": "duplicate of row 1 in this file"}
    ]


def test_enum_values_are_normalized_and_validated_with_readable_messages(
    client: TestClient,
) -> None:
    content = _csv(
        "part_number,description,part_type,qualification_status,lifecycle_status,source_type",
        "PV-1001,Valve,valve,In Qualification,Active,Vendor",
        "PV-1002,Valve,valve,approved,activ,internal",
    )
    rows = _import_parts(client, content).json()["rows"]
    assert rows[0]["action"] == "create"
    assert rows[0]["changes"]["qualification_status"] == [None, "in_qualification"]
    assert rows[0]["changes"]["source_type"] == [None, "vendor"]
    errors = {error["field"]: error["message"] for error in rows[1]["errors"]}
    assert errors["qualification_status"] == (
        "must be one of: disqualified, in_qualification, qualified, unqualified (got 'approved')"
    )
    assert errors["lifecycle_status"].startswith("must be one of: active, draft, legacy")
    assert errors["lifecycle_status"].endswith("(got 'activ')")


def test_import_size_limit(client: TestClient) -> None:
    content = b"part_number\n" + b"x" * MAX_IMPORT_BYTES
    response = _import_parts(client, content)
    assert response.status_code == 413
    pasted = client.post(
        "/parts/import", json={"headers": ["part_number"], "rows": [["x" * MAX_IMPORT_BYTES]]}
    )
    assert pasted.status_code == 413


def test_viewer_cannot_import_but_can_download_templates(client: TestClient) -> None:
    project_id = _project(client)
    viewer = _viewer(client)
    content = _csv("part_number,description,part_type", "PV-1001,Valve,valve")
    assert _import_parts(viewer, content).status_code == 403
    requirement_import = viewer.post(
        f"/projects/{project_id}/requirements/import",
        json={"headers": ["key"], "rows": [["R-1"]]},
    )
    assert requirement_import.status_code == 403
    assert viewer.get("/parts/import-template").status_code == 200
    assert client.get("/parts").json() == []


def test_requirement_import_from_pasted_rows(client: TestClient) -> None:
    project_id = _project(client)
    other_id = _project(client, "Other")
    client.post(
        "/requirements",
        json={
            "project_id": other_id,
            "key": "REQ-1",
            "title": "Elsewhere",
            "text": "Other project",
            "requirement_type": "safety",
        },
    )
    path = f"/projects/{project_id}/requirements/import"
    body = {
        "rows": [
            ["ID", "Title", "Description", "Type", "Verification Method", "Constraint", "Notes"],
            ["REQ-1", "Wetted", "316L wetted parts", "materials", "inspection", "", "x"],
            [
                "REQ-2",
                "Rating",
                "Rated parts",
                "safety",
                "analysis",
                '{"kind": "pressure_rating_min", "values": ["100"]}',
                "",
            ],
            ["REQ-3", "Broken", "Bad constraint", "safety", "", "{not json", ""],
        ]
    }
    report = client.post(path, json=body).json()
    assert report["entity"] == "requirement"
    assert report["mapping"]["Notes"] is None
    # REQ-1 of another project is not a match: requirement keys are per project.
    assert [row["action"] for row in report["rows"]] == ["create", "create", "error"]
    assert report["rows"][1]["changes"]["constraint"][1]["kind"] == "pressure_rating_min"
    assert report["rows"][2]["errors"][0]["field"] == "constraint"

    body["rows"].pop()
    committed = client.post(path, params={"dry_run": "false"}, json=body)
    assert committed.status_code == 200, committed.text
    requirements = client.get(f"/projects/{project_id}/requirements").json()
    assert sorted(requirement["key"] for requirement in requirements) == ["REQ-1", "REQ-2"]
    assert all(requirement["project_id"] == project_id for requirement in requirements)

    body = {"rows": [{"Key": "req-1", "Status": "approved"}, {"Key": "REQ-2", "Owner": "QA"}]}
    upsert = client.post(path, params={"mode": "upsert", "dry_run": "false"}, json=body).json()
    assert upsert["summary"] == {"create": 0, "update": 2, "unchanged": 0, "error": 0}
    statuses = {req["key"]: req["status"] for req in client.get(path[: -len("/import")]).json()}
    assert statuses == {"REQ-1": "approved", "REQ-2": "draft"}
    project_changes = [
        change
        for change in client.get("/changes").json()
        if change["object_type"] == "project" and change["object_id"] == project_id
    ]
    assert any("Imported requirements from pasted rows" in c["summary"] for c in project_changes)


def _document(item_id: str) -> dict:
    return {
        "schemaVersion": 1,
        "sheet": {
            "size": "A3",
            "orientation": "landscape",
            "frame": {"kind": "basic", "columns": 4, "rows": 3, "margin": 10},
        },
        "layers": [],
        "items": [
            {
                "id": item_id,
                "kind": "symbol",
                "layer": "symbols",
                "symbol": {"library": "fsdp", "key": "valve", "version": 1},
                "position": {"x": 50, "y": 50},
                "rotation": 0,
                "tag": item_id.upper(),
                "fields": {},
            }
        ],
        "meta": {"grid": 2.5},
    }


def _indexed_sheet(client: TestClient, project_id: str, part_id: str | None = None) -> str:
    drawing = client.post(
        f"/projects/{project_id}/drawings",
        json={"title": "P&ID", "first_sheet": {"document": _document("hv")}},
    ).json()
    sheet_id = drawing["sheets"][0]["id"]
    saved = client.put(
        f"/sheets/{sheet_id}",
        json={
            "document": _document("hv"),
            "index": {
                "items": [{"item_id": "hv", "kind": "symbol", "tag": "HV", "part_id": part_id}],
                "lines": [],
            },
            "drc": {"findings": [], "checks": []},
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["index_stale"] is False
    return sheet_id


def test_requirement_import_marks_sheets_stale_only_for_constraints(client: TestClient) -> None:
    project_id = _project(client)
    sheet_id = _indexed_sheet(client, project_id)
    path = f"/projects/{project_id}/requirements/import"
    manual = {"rows": [["key", "title", "text", "requirement_type"], ["R-1", "T", "X", "s"]]}
    assert client.post(path, params={"dry_run": "false"}, json=manual).status_code == 200
    assert client.get(f"/sheets/{sheet_id}").json()["index_stale"] is False

    checked = {
        "headers": ["key", "title", "text", "requirement_type", "constraint"],
        "rows": [["R-2", "T", "X", "s", '{"kind": "material_in", "values": ["316L"]}']],
    }
    assert client.post(path, params={"dry_run": "false"}, json=checked).status_code == 200
    assert client.get(f"/sheets/{sheet_id}").json()["index_stale"] is True


def test_part_upsert_marks_sheets_using_the_part_stale(client: TestClient) -> None:
    project_id = _project(client)
    part = _part(client, "PV-1001")
    sheet_id = _indexed_sheet(client, project_id, part["id"])
    unchanged = _csv("part_number,material", "PV-1001,316L")
    _import_parts(client, unchanged, mode="upsert", dry_run="false")
    assert client.get(f"/sheets/{sheet_id}").json()["index_stale"] is False
    changed = _csv("part_number,material", "PV-1001,Brass")
    _import_parts(client, changed, mode="upsert", dry_run="false")
    assert client.get(f"/sheets/{sheet_id}").json()["index_stale"] is True


def test_import_templates_download_and_round_trip(client: TestClient) -> None:
    project_id = _project(client)
    csv_template = client.get("/parts/import-template", params={"format": "csv"})
    assert csv_template.status_code == 200
    assert csv_template.headers["content-type"].startswith("text/csv")
    assert "parts-import-template.csv" in csv_template.headers["content-disposition"]
    assert csv_template.content.startswith(b"\xef\xbb\xbf")
    header, example = list(csv.reader(io.StringIO(csv_template.content.decode("utf-8-sig"))))
    assert header[:3] == ["part_number", "revision", "description"]
    assert "lifecycle_status" in header and len(example) == len(header)

    xlsx_template = client.get(
        f"/projects/{project_id}/requirements/import-template", params={"format": "xlsx"}
    )
    assert xlsx_template.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx_template.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "key",
        "title",
        "text",
        "requirement_type",
        "verification_method",
        "status",
        "owner",
        "constraint",
    )
    assert rows[1][0] == "REQ-001"

    # The templates' example rows are valid imports.
    part_report = _import_parts(client, csv_template.content).json()
    assert part_report["summary"] == {"create": 1, "update": 0, "unchanged": 0, "error": 0}
    requirement_report = _upload(
        client,
        f"/projects/{project_id}/requirements/import",
        xlsx_template.content,
        "requirements.xlsx",
    ).json()
    assert requirement_report["summary"]["create"] == 1, requirement_report

    assert client.get("/parts/import-template", params={"format": "pdf"}).status_code == 422
    assert client.get("/projects/missing/requirements/import-template").status_code == 404


def test_line_classes_accept_an_xlsx_upload(client: TestClient) -> None:
    project_id = _project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Material", "Rating", "Sizes"])
    sheet.append(["CS150", "Carbon steel", "150#", "1/2;3/4;1"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = client.post(
        f"/projects/{project_id}/line-classes/import-file",
        files={"file": ("classes.xlsx", buffer.getvalue(), "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    assert response.json() == {"created": 1, "updated": 0, "errors": []}
    classes = client.get(f"/projects/{project_id}/line-classes").json()
    assert [(row["name"], row["material"], row["sizes"]) for row in classes] == [
        ("CS150", "Carbon steel", ["1/2", "3/4", "1"])
    ]
