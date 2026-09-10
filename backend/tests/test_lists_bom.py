"""Phase 4: sheet index on save, engineering lists, drawing BoM, part usage."""

import csv
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _project(client: TestClient) -> str:
    return client.post("/projects", json={"name": "AMB2", "part_name_prefix": "AMB2"}).json()["id"]


def _drawing(client: TestClient, project_id: str) -> dict:
    return client.post(
        f"/projects/{project_id}/drawings",
        json={"title": "P&ID\nHELIUM FILL", "size": "ANSI_E", "units": "in"},
    ).json()


def _part(client: TestClient, number: str, **extra) -> dict:
    body = {
        "part_number": number,
        "description": f"Part {number}",
        "part_type": "valve",
        "material": "316L",
        "pressure_rating_bar": 200,
        "manufacturer": "Swagelok",
        "qualification_status": "qualified",
        **extra,
    }
    return client.post("/parts", json=body).json()


def _index(part_id: str | None = None) -> dict:
    return {
        "items": [
            {
                "item_id": "hv",
                "kind": "symbol",
                "category": "valve",
                "symbol_key": "ball_valve",
                "symbol_name": "Ball valve",
                "tag": "HV-3201",
                "zone": "D-4",
                "x": 100,
                "y": 120,
                "part_id": part_id,
                "fields": {"size": '1/4"', "service": "GHe", "line_number": "3101"},
            },
            {
                "item_id": "hv2",
                "kind": "symbol",
                "category": "valve",
                "symbol_key": "ball_valve",
                "symbol_name": "Ball valve",
                "tag": "HV-3202",
                "zone": "D-3",
                "part_id": part_id,
                "dnp": True,
                "fields": {"size": '1/4"'},
            },
            {
                "item_id": "pt",
                "kind": "symbol",
                "category": "instrument",
                "symbol_key": "instrument_field",
                "symbol_name": "Field instrument",
                "tag": "PT-3204",
                "zone": "C-3",
                "spare": 1,
                "fields": {"mounting": "Field"},
            },
            {
                "item_id": "vent",
                "kind": "symbol",
                "category": "connector",
                "symbol_key": "off_page_connector",
                "symbol_name": "Off-page connector",
                "label": "TO VENT",
                "zone": "A-1",
                "fields": {"ref": "A", "target": "SHT 2 / D-4"},
            },
            {
                "item_id": "vc",
                "kind": "equipment",
                "symbol_name": "Equipment boundary",
                "tag": "VC-1",
                "label": "VACUUM CHAMBER",
                "zone": "B-2",
                "fields": {"nozzle_count": 2},
            },
            {
                "item_id": "tank",
                "kind": "equipment",
                "symbol_name": "Equipment boundary",
                "tag": "N-4200",
                "label": "STORAGE TANK",
                "zone": "B-3",
                "fields": {"nozzle_count": 1},
            },
        ],
        "lines": [
            {
                "line_id": "l1",
                "line_number": "3101",
                "line_type": "process",
                "service": "GHe",
                "size": '1/4"',
                "spec": 'ST x .035" WALL',
                "line_class": "HE-1",
                "from_tag": "HV-3201",
                "to_tag": "PT-3204",
                "zone": "D-4",
                "length_mm": 120,
                "length_m": 1.5,
                "connection_count": 2,
                "tee_count": 1,
                "design_pressure": "2500 psig",
            },
            {
                "line_id": "l2",
                "line_number": "3102",
                "line_type": "process",
                "service": "GHe",
                "size": '1/4"',
                "line_class": "HE-1",
                "zone": "C-4",
                "length_mm": 40,
                "length_m": 0.5,
                "connection_count": 1,
            },
            {
                "line_id": "s1",
                "line_type": "signal_electric",
                "zone": "C-3",
                "length_mm": 30,
            },
        ],
    }


def test_sheet_index_is_replaced_on_save(client: TestClient) -> None:
    project_id = _project(client)
    drawing = _drawing(client, project_id)
    sheet_id = drawing["sheets"][0]["id"]
    part = _part(client, "AMB2-001")
    saved = client.put(f"/sheets/{sheet_id}", json={"index": _index(part["id"])})
    assert saved.status_code == 200, saved.text
    index = client.get(f"/sheets/{sheet_id}/index").json()
    assert {item["item_id"] for item in index["items"]} == {"hv", "hv2", "pt", "vent", "vc", "tank"}
    assert index["items"][0]["part_id"] == part["id"]
    assert len(index["lines"]) == 3
    # Saving a smaller index replaces the rows; an unknown part id is dropped, not rejected.
    smaller = {"items": [{"item_id": "hv", "kind": "symbol", "part_id": "missing"}], "lines": []}
    client.put(f"/sheets/{sheet_id}", json={"index": smaller})
    index = client.get(f"/sheets/{sheet_id}/index").json()
    assert [item["item_id"] for item in index["items"]] == ["hv"]
    assert index["items"][0]["part_id"] is None
    assert index["lines"] == []


def test_lists_per_drawing_and_project_in_json_csv_and_xlsx(client: TestClient) -> None:
    project_id = _project(client)
    drawing = _drawing(client, project_id)
    sheet_id = drawing["sheets"][0]["id"]
    part = _part(client, "AMB2-001")
    client.put(f"/sheets/{sheet_id}", json={"index": _index(part["id"])})

    instruments = client.get(f"/drawings/{drawing['id']}/lists/instrument").json()
    assert instruments["title"] == "Instrument index"
    assert instruments["header"]["drawing_number"] == drawing["number"]
    assert instruments["header"]["revision"] == "-"
    assert [row["tag"] for row in instruments["rows"]] == ["PT-3204"]
    assert instruments["rows"][0]["zone"] == "C-3"
    assert instruments["rows"][0]["sheet_no"] == 1
    assert [column["key"] for column in instruments["columns"]][-2:] == ["sheet_no", "zone"]

    valves = client.get(f"/drawings/{drawing['id']}/lists/valve").json()
    assert [(row["tag"], row["dnp"], row["part_number"]) for row in valves["rows"]] == [
        ("HV-3201", "", "AMB2-001"),
        ("HV-3202", "DNP", "AMB2-001"),
    ]
    assert valves["rows"][0]["line_number"] == "3101"

    lines = client.get(f"/drawings/{drawing['id']}/lists/line").json()
    assert [row["line_number"] for row in lines["rows"]] == [None, "3101", "3102"] or [
        row["line_number"] for row in lines["rows"]
    ] == ["3101", "3102", None]
    numbered = [row for row in lines["rows"] if row["line_number"] == "3101"][0]
    assert numbered["from_tag"] == "HV-3201"
    assert numbered["to_tag"] == "PT-3204"
    assert numbered["length_m"] == 1.5

    equipment = client.get(f"/drawings/{drawing['id']}/lists/equipment").json()
    assert [(row["tag"], row["name"], row["nozzle_count"]) for row in equipment["rows"]] == [
        ("N-4200", "STORAGE TANK", 1),
        ("VC-1", "VACUUM CHAMBER", 2),
    ]
    tie_ins = client.get(f"/drawings/{drawing['id']}/lists/tie-in").json()
    assert tie_ins["rows"][0]["target"] == "SHT 2 / D-4"
    assert tie_ins["rows"][0]["tag"] == "TO VENT"

    assert client.get(f"/drawings/{drawing['id']}/lists/bogus").status_code == 404
    assert client.get(f"/drawings/{drawing['id']}/lists/valve?format=pdf").status_code == 400

    csv_response = client.get(f"/drawings/{drawing['id']}/lists/valve?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert f"{drawing['number']}-valve-list.csv" in csv_response.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(csv_response.text)))
    assert rows[0] == ["List", "Valve list"]
    assert ["Drawing Number", drawing["number"]] in rows
    header_index = rows.index(
        [
            "Tag",
            "Type",
            "Actuator",
            "Size",
            "Service",
            "Line",
            "Part",
            "DNP",
            "Notes",
            "Sheet",
            "Zone",
        ]
    )
    assert rows[header_index + 1][0] == "HV-3201"
    assert rows[header_index + 1][-1] == "D-4"

    xlsx_response = client.get(f"/drawings/{drawing['id']}/lists/line?format=xlsx")
    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"].endswith("spreadsheetml.sheet")
    workbook = load_workbook(io.BytesIO(xlsx_response.content))
    sheet = workbook.active
    assert sheet.title == "Line list"
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert ["Drawing Number", drawing["number"]] in [row[:2] for row in values]
    header_row = next(row for row in values if row[0] == "Line")
    assert header_row[-1] == "Zone"
    assert any(row[0] == "3101" and row[-1] == "D-4" for row in values)

    # Project scope prefixes a drawing column and spans every drawing.
    second = _drawing(client, project_id)
    client.put(
        f"/sheets/{second['sheets'][0]['id']}",
        json={
            "index": {
                "items": [
                    {"item_id": "x", "kind": "symbol", "category": "valve", "tag": "HV-4201"}
                ],
                "lines": [],
            }
        },
    )
    project_valves = client.get(f"/projects/{project_id}/lists/valve").json()
    assert project_valves["scope"] == "project"
    assert project_valves["columns"][0]["key"] == "drawing_number"
    assert [(row["drawing_number"], row["tag"]) for row in project_valves["rows"]] == [
        (drawing["number"], "HV-3201"),
        (drawing["number"], "HV-3202"),
        (second["number"], "HV-4201"),
    ]
    assert "drawing_number" not in project_valves["header"]


def test_drawing_bom_rolls_up_parts_dnp_spares_and_bulk_items(client: TestClient) -> None:
    project_id = _project(client)
    drawing = _drawing(client, project_id)
    sheet_id = drawing["sheets"][0]["id"]
    part = _part(client, "AMB2-001")
    client.put(f"/sheets/{sheet_id}", json={"index": _index(part["id"])})

    created = client.post(f"/drawings/{drawing['id']}/bom")
    assert created.status_code == 201, created.text
    snapshot = created.json()
    assert snapshot["drawing_id"] == drawing["id"]
    assert snapshot["diagram_id"] is None
    assert snapshot["source_kind"] == "drawing"
    assert snapshot["revision"] == 1
    rows = {row["description"]: row for row in snapshot["rows"]}

    valve = rows["Part AMB2-001"]
    assert valve["kind"] == "part"
    assert valve["quantity"] == 1  # HV-3202 is DNP
    assert valve["component_tags"] == ["HV-3201"]
    assert valve["dnp_tags"] == ["HV-3202"]
    assert valve["sheets"] == [1]

    instrument = rows["Field instrument"]
    assert instrument["kind"] == "unassigned"
    assert instrument["part_number"] is None
    assert instrument["quantity"] == 1
    assert instrument["spare_quantity"] == 1

    chamber = rows["VACUUM CHAMBER"]
    assert chamber["component_tags"] == ["VC-1"]
    assert rows["STORAGE TANK"]["component_tags"] == ["N-4200"]
    assert "Off-page connector" not in rows

    tube = rows['Tube 1/4" HE-1']
    assert tube["kind"] == "bulk"
    assert tube["unit"] == "m"
    assert tube["quantity"] == 2.0
    assert tube["component_tags"] == ["3101", "3102"]
    assert rows['Tube fitting 1/4"']["quantity"] == 3
    assert rows['Tube tee 1/4"']["quantity"] == 1
    # Signal lines are not bulk material.
    assert all("signal" not in description.lower() for description in rows)

    again = client.post(f"/drawings/{drawing['id']}/bom").json()
    assert again["revision"] == 2
    listed = client.get(f"/drawings/{drawing['id']}/bom").json()
    assert [entry["revision"] for entry in listed] == [2, 1]

    project_boms = client.get(f"/projects/{project_id}/bom").json()
    assert {entry["diagram_name"] for entry in project_boms} == {drawing["number"]}
    assert project_boms[0]["source_kind"] == "drawing"

    readiness = client.get(f"/bom/{snapshot['id']}/readiness").json()
    codes = {(issue["code"], issue["severity"]) for issue in readiness["issues"]}
    assert ("no_part", "blocking") in codes
    assert readiness["blocking_count"] >= 1
    assert readiness["ready"] is False
    unassigned = next(issue for issue in readiness["issues"] if issue["code"] == "no_part")
    assert "PT-3204" in unassigned["component_tags"] or "VC-1" in unassigned["component_tags"]

    csv_export = client.get(f"/bom/{snapshot['id']}/csv")
    assert csv_export.status_code == 200
    assert "AMB2-001" in csv_export.text


def test_bulk_rows_without_size_or_class_are_warnings(client: TestClient) -> None:
    project_id = _project(client)
    drawing = _drawing(client, project_id)
    sheet_id = drawing["sheets"][0]["id"]
    client.put(
        f"/sheets/{sheet_id}",
        json={
            "index": {
                "items": [],
                "lines": [
                    {
                        "line_id": "a",
                        "line_type": "process",
                        "length_m": 1.0,
                        "connection_count": 1,
                    },
                    {"line_id": "b", "line_type": "process", "size": '1/2"', "length_m": 2.0},
                ],
            }
        },
    )
    snapshot = client.post(f"/drawings/{drawing['id']}/bom").json()
    readiness = client.get(f"/bom/{snapshot['id']}/readiness").json()
    codes = sorted((issue["code"], issue["severity"]) for issue in readiness["issues"])
    assert codes == [
        ("line_no_class", "warning"),
        ("line_no_size", "warning"),
        ("line_no_size", "warning"),
    ]
    assert readiness["blocking_count"] == 0
    assert readiness["warning_count"] == 3


def test_part_usage_and_delete_guard_include_drawings(client: TestClient) -> None:
    project_id = _project(client)
    drawing = _drawing(client, project_id)
    sheet_id = drawing["sheets"][0]["id"]
    part = _part(client, "AMB2-001")
    client.put(f"/sheets/{sheet_id}", json={"index": _index(part["id"])})
    snapshot = client.post(f"/drawings/{drawing['id']}/bom").json()

    usage = client.get(f"/parts/{part['id']}/usage").json()
    assert usage["components"] == []
    assert [(item["tag"], item["zone"], item["dnp"]) for item in usage["drawing_items"]] == [
        ("HV-3201", "D-4", False),
        ("HV-3202", "D-3", True),
    ]
    assert usage["drawing_items"][0]["drawing_number"] == drawing["number"]
    assert usage["drawing_items"][0]["sheet_no"] == 1
    assert [entry["id"] for entry in usage["bom_snapshots"]] == [snapshot["id"]]

    blocked = client.delete(f"/parts/{part['id']}")
    assert blocked.status_code == 409
    assert "2 component instance(s)" in blocked.json()["detail"]

    # Deleting the drawing cascades the index rows and frees the part.
    assert client.delete(f"/drawings/{drawing['id']}").status_code == 204
    assert client.get(f"/parts/{part['id']}/usage").json()["drawing_items"] == []
    assert client.delete(f"/parts/{part['id']}").status_code == 204
