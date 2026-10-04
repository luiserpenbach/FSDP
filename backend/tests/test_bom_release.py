"""Drawing BoM release gate, immutability, XLSX export, and query batching."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.db import get_db
from app.main import app


def _project(client: TestClient) -> str:
    return client.post("/projects", json={"name": "AMB2", "part_name_prefix": "AMB2"}).json()["id"]


def _part(client: TestClient, number: str, **extra) -> dict:
    body = {
        "part_number": number,
        "description": f"Valve {number}",
        "part_type": "valve",
        "material": "316L",
        "pressure_rating_bar": 200,
        "manufacturer": "Swagelok",
        "qualification_status": "qualified",
        **extra,
    }
    return client.post("/parts", json=body).json()


def _index(*part_ids: str | None) -> dict:
    return {
        "items": [
            {
                "item_id": f"hv{n}",
                "kind": "symbol",
                "category": "valve",
                "symbol_key": "ball_valve",
                "symbol_name": "Ball valve",
                "tag": f"HV-{n}",
                "part_id": part_id,
            }
            for n, part_id in enumerate(part_ids, start=1)
        ],
        "lines": [],
    }


def _drawing_with_bom(client: TestClient, *part_ids: str | None, drc: bool = True) -> dict:
    project_id = _project(client)
    drawing = client.post(f"/projects/{project_id}/drawings", json={"title": "P&ID"}).json()
    body: dict = {"index": _index(*part_ids)}
    if drc:
        body["drc"] = {"findings": [], "checks": []}
    client.put(f"/sheets/{drawing['sheets'][0]['id']}", json=body)
    snapshot = client.post(f"/drawings/{drawing['id']}/bom")
    assert snapshot.status_code == 201, snapshot.text
    return {"drawing": drawing, "snapshot": snapshot.json(), "project_id": project_id}


def test_release_refused_with_blocking_issues(client: TestClient) -> None:
    good = _part(client, "AMB2-001")
    setup = _drawing_with_bom(client, good["id"], None)
    snapshot = setup["snapshot"]

    refused = client.put(f"/bom/{snapshot['id']}/status", json={"status": "released"})
    assert refused.status_code == 409
    detail = refused.json()["detail"]
    assert "blocking" in detail["message"]
    assert [issue["code"] for issue in detail["issues"]] == ["no_part"]
    assert client.get(f"/bom/{snapshot['id']}/readiness").json()["blocking_count"] == 1
    assert client.get(f"/drawings/{setup['drawing']['id']}/bom").json()[0]["status"] == "draft"


def test_stale_sheets_block_bom_release(client: TestClient) -> None:
    good = _part(client, "AMB2-001")
    setup = _drawing_with_bom(client, good["id"], drc=False)
    snapshot = setup["snapshot"]
    sheet_id = setup["drawing"]["sheets"][0]["id"]
    assert [entry["sheet_id"] for entry in snapshot["stale_sheets"]] == [sheet_id]
    assert snapshot["drawing_revision"] == "-"

    readiness = client.get(f"/bom/{snapshot['id']}/readiness").json()
    assert [(issue["code"], issue["severity"]) for issue in readiness["issues"]] == [
        ("stale_index", "blocking")
    ]
    refused = client.put(f"/bom/{snapshot['id']}/status", json={"status": "released"})
    assert refused.status_code == 409
    assert refused.json()["detail"]["issues"][0]["code"] == "stale_index"


def test_released_bom_is_immutable(client: TestClient) -> None:
    good = _part(client, "AMB2-001")
    setup = _drawing_with_bom(client, good["id"])
    snapshot = setup["snapshot"]
    assert snapshot["stale_sheets"] == []

    released = client.put(f"/bom/{snapshot['id']}/status", json={"status": "released"})
    assert released.status_code == 200, released.text
    assert released.json()["status"] == "released"
    assert released.json()["released_by"] == "engineer@fsdp.test"
    assert released.json()["released_at"]

    again = client.put(f"/bom/{snapshot['id']}/status", json={"status": "released"})
    assert again.status_code == 200
    back = client.put(f"/bom/{snapshot['id']}/status", json={"status": "draft"})
    assert back.status_code == 409
    assert "immutable" in back.json()["detail"]
    unknown = client.put(f"/bom/{snapshot['id']}/status", json={"status": "approved"})
    assert unknown.status_code == 422

    history = client.get(f"/drawings/{setup['drawing']['id']}/bom").json()
    assert [(entry["revision"], entry["status"], entry["released_by"]) for entry in history] == [
        (1, "released", "engineer@fsdp.test")
    ]
    client.post(f"/drawings/{setup['drawing']['id']}/bom")
    history = client.get(f"/drawings/{setup['drawing']['id']}/bom").json()
    assert [(entry["revision"], entry["status"]) for entry in history] == [
        (2, "draft"),
        (1, "released"),
    ]


def test_bom_xlsx_export_has_header_block(client: TestClient) -> None:
    good = _part(client, "=AMB2-001")
    setup = _drawing_with_bom(client, good["id"])
    snapshot = setup["snapshot"]
    drawing = setup["drawing"]

    response = client.get(f"/bom/{snapshot['id']}/xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")
    assert f"bom-{drawing['number'].lower()}-rev1.xlsx" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    header = {row[0]: row[1] for row in values[:10] if row[0]}
    assert header["List"] == "Bill of materials"
    assert header["Project"] == "AMB2"
    assert header["Drawing Number"] == drawing["number"]
    # spreadsheet_safe quotes a leading "-" like any formula prefix.
    assert header["Drawing Revision"] in {"-", "'-"}
    assert header["Bom Revision"] == "1"
    assert header["Status"] == "draft"
    assert header["Generated"]
    columns = next(row for row in values if row[0] == "Part number")
    assert "Quantity" in columns
    data = values[values.index(columns) + 1]
    assert data[0] == "'=AMB2-001"
    assert data[columns.index("Quantity")] == 1
    assert data[columns.index("Tags")] == "HV-1"


def test_legacy_diagram_bom_xlsx(client: TestClient) -> None:
    project = client.post("/projects", json={"name": "Legacy"}).json()
    system = client.post(f"/projects/{project['id']}/systems", json={"name": "Feed"}).json()
    diagram = client.post(f"/systems/{system['id']}/diagrams", json={"name": "P&ID"}).json()
    client.post(f"/diagrams/{diagram['id']}/components", json={"tag": "V-1"})
    snapshot = client.post(f"/diagrams/{diagram['id']}/bom").json()
    sheet = load_workbook(io.BytesIO(client.get(f"/bom/{snapshot['id']}/xlsx").content)).active
    header = {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}
    assert header["Project"] == "Legacy"
    assert header["Diagram"] == "P&ID"


def test_readiness_loads_parts_in_one_query(client: TestClient) -> None:
    parts = [_part(client, f"AMB2-00{n}") for n in range(1, 6)]
    setup = _drawing_with_bom(client, *(part["id"] for part in parts))

    db = next(app.dependency_overrides[get_db]())
    statements: list[str] = []

    def count(_conn, _cursor, statement, *_args) -> None:
        statements.append(statement)

    engine = db.get_bind()
    event.listen(engine, "before_cursor_execute", count)
    try:
        readiness = client.get(f"/bom/{setup['snapshot']['id']}/readiness").json()
    finally:
        event.remove(engine, "before_cursor_execute", count)
        db.close()
    assert readiness["row_count"] == 5
    part_queries = [sql for sql in statements if "FROM parts" in sql]
    assert len(part_queries) == 1
