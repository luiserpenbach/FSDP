"""XLSX export of a table as the DataGrid shows it (POST /exports/xlsx)."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _workbook(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return load_workbook(io.BytesIO(response.content))


def _table(sheet) -> list[list]:
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    # Title block, blank row, then the header row.
    header_index = next(i for i, row in enumerate(values) if row and row[0] == "Part number")
    return values[header_index:]


def test_exports_visible_columns_and_rows_in_order(client: TestClient) -> None:
    response = client.post(
        "/exports/xlsx",
        json={
            "title": "Parts catalog",
            "file_name": "parts",
            "columns": [
                {"key": "part_number", "header": "Part number"},
                {"key": "pressure_rating_bar", "header": "Pressure (bar)"},
                {"key": "preferred", "header": "Preferred"},
                {"key": "notes", "header": "Notes"},
            ],
            "rows": [
                {
                    "part_number": "PV-2",
                    "pressure_rating_bar": 206.5,
                    "preferred": True,
                    "notes": None,
                },
                {"part_number": "PV-10", "pressure_rating_bar": None, "preferred": False},
            ],
        },
    )
    assert response.headers["content-disposition"] == 'attachment; filename="parts.xlsx"'
    sheet = _workbook(response).active
    assert sheet.title == "Parts catalog"
    assert sheet["A1"].value == "List"
    assert sheet["B1"].value == "Parts catalog"
    assert _table(sheet) == [
        ["Part number", "Pressure (bar)", "Preferred", "Notes"],
        ["PV-2", 206.5, True, None],
        ["PV-10", None, False, None],
    ]


def test_guards_formula_injection_and_illegal_characters(client: TestClient) -> None:
    response = client.post(
        "/exports/xlsx",
        json={
            "title": "Requirements: Q3/Q4 [draft]",
            "file_name": "../req list",
            "columns": [{"key": "a", "header": "Part number"}, {"key": "b", "header": "Text"}],
            "rows": [{"a": '=HYPERLINK("http://x")', "b": "bell\x07 -5 kg"}],
        },
    )
    assert response.headers["content-disposition"] == 'attachment; filename="req-list.xlsx"'
    sheet = _workbook(response).active
    assert sheet.title == "Requirements Q3 Q4 draft"
    rows = _table(sheet)
    assert rows[1] == ['\'=HYPERLINK("http://x")', "bell -5 kg"]


def test_rejects_requests_without_columns(client: TestClient) -> None:
    response = client.post("/exports/xlsx", json={"columns": [], "rows": []})
    assert response.status_code == 422


def test_requires_a_signed_in_user(client: TestClient) -> None:
    client.post("/auth/logout")
    response = client.post(
        "/exports/xlsx", json={"columns": [{"key": "a", "header": "A"}], "rows": []}
    )
    assert response.status_code == 401
