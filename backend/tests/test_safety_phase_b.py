"""Safety phase B: failure-mode library, FMEA worksheets generated from the
sheet index, staleness on save, the release gate, releases, diffs, exports,
row controls, and comments."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader


def _project(client: TestClient) -> str:
    return client.post("/projects", json={"name": "LOX GSE", "part_name_prefix": "LOX"}).json()[
        "id"
    ]


def _item(item_id: str, category: str, symbol_key: str, tag: str, **fields: object) -> dict:
    return {
        "item_id": item_id,
        "kind": "symbol",
        "category": category,
        "symbol_key": symbol_key,
        "symbol_name": symbol_key.replace("_", " ").title(),
        "tag": tag,
        "label": None,
        "zone": "C-4",
        "x": 10,
        "y": 10,
        "part_id": fields.pop("part_id", None),
        "dnp": bool(fields.pop("dnp", False)),
        "spare": 0,
        "fields": {"service": "LOX", "line_number": "L-2014", **fields},
    }


LOX_ITEMS = [
    _item("fv201", "valve", "pneumatic_valve", "FV-201"),
    _item("cv203", "valve", "check_valve", "CV-203"),
    _item("trv201", "valve", "relief_valve", "TRV-201"),
    _item("f210", "inline", "filter", "F-210"),
    _item("pt205", "instrument", "instrument", "PT-205"),
    _item("zs201", "instrument", "instrument", "ZS-201", line_number="L-2014"),
    _item("qd201", "inline", "quick_disconnect", "QD-201", line_number="L-2015"),
    _item("spare", "valve", "ball_valve", "HV-9", dnp=True),
]


def _drawing_with_items(client: TestClient, project_id: str, items: list[dict]) -> tuple[str, str]:
    drawing = client.post(
        f"/projects/{project_id}/drawings", json={"title": "LOX P&ID", "size": "ANSI_E"}
    ).json()
    sheet_id = drawing["sheets"][0]["id"]
    saved = client.put(f"/sheets/{sheet_id}", json={"index": {"items": items, "lines": []}})
    assert saved.status_code == 200, saved.text
    return drawing["id"], sheet_id


def _worksheet(client: TestClient, project_id: str, drawing_id: str) -> dict:
    response = client.post(
        f"/projects/{project_id}/fmea",
        json={
            "title": "LOX fill and drain",
            "drawing_id": drawing_id,
            "operating_modes": ["fast_fill", "hold"],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_failure_mode_library_seeds_and_resolves(client: TestClient) -> None:
    entries = client.get("/failure-modes").json()
    assert len(entries) > 40
    names = {(entry["category"], entry["symbol_key"], entry["name"]) for entry in entries}
    assert ("valve", None, "fails_open") in names
    assert ("valve", "check_valve", "reverse_flow") in names
    assert ("inline", "quick_disconnect", "premature_disconnect") in names
    # Seeding is idempotent.
    assert len(client.get("/failure-modes").json()) == len(entries)
    created = client.post(
        "/failure-modes",
        json={
            "category": "valve",
            "symbol_key": "cryo_valve",
            "name": "ice_lock",
            "title": "Ice lock",
        },
    )
    assert created.status_code == 201, created.text
    duplicate = client.post(
        "/failure-modes",
        json={
            "category": "valve",
            "symbol_key": "cryo_valve",
            "name": "ice_lock",
            "title": "Ice lock",
        },
    )
    assert duplicate.status_code == 409


def test_generate_rows_from_the_sheet_index(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, sheet_id = _drawing_with_items(client, project_id, LOX_ITEMS)
    worksheet = _worksheet(client, project_id, drawing_id)
    assert worksheet["row_count"] == 0

    generated = client.post(f"/fmea/{worksheet['id']}/generate", json={})
    assert generated.status_code == 200, generated.text
    body = generated.json()
    assert body["kept"] == 0 and body["stale"] == 0
    rows = client.get(f"/fmea/{worksheet['id']}/rows").json()
    assert body["added"] == len(rows) > 20
    by_item: dict[str, list[dict]] = {}
    for row in rows:
        by_item.setdefault(row["item_tag"], []).append(row)
    # DNP items are skipped; check valves replace the generic valve modes.
    assert "HV-9" not in by_item
    assert {row["failure_mode_name"] for row in by_item["CV-203"]} == {
        "reverse_flow",
        "stuck_closed",
        "chatter",
    }
    # Actuated valves get the generic set plus the actuator mode.
    fv_modes = {row["failure_mode_name"] for row in by_item["FV-201"]}
    assert {"fails_open", "fails_closed", "loss_of_supply"} <= fv_modes
    # Effect templates are rendered and detection comes from an instrument on the line.
    fails_closed = next(
        row for row in by_item["FV-201"] if row["failure_mode_name"] == "fails_closed"
    )
    assert fails_closed["local_effect"] == "FV-201 blocks LOX; no flow downstream"
    assert fails_closed["detected_by_tag"] in {"PT-205", "ZS-201"}
    assert fails_closed["detection_kind"] == "instrument"
    assert fails_closed["operating_modes"] == ["fast_fill", "hold"]
    assert fails_closed["severity"] == 6
    assert fails_closed["item_zone"] == "C-4"
    assert fails_closed["sheet_no"] == 1

    # Idempotent: a second generate keeps every row.
    again = client.post(f"/fmea/{worksheet['id']}/generate", json={}).json()
    assert again["added"] == 0 and again["kept"] == len(rows)
    view = client.get(f"/fmea/{worksheet['id']}").json()
    assert view["row_count"] == len(rows)
    assert view["drawing_number"].startswith("LOX")


def test_rows_go_stale_when_the_sheet_changes(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, sheet_id = _drawing_with_items(client, project_id, LOX_ITEMS)
    worksheet = _worksheet(client, project_id, drawing_id)
    client.post(f"/fmea/{worksheet['id']}/generate", json={"categories": ["valve"]})
    part = client.post(
        "/parts",
        json={
            "part_number": "AMPH-VL-022",
            "description": "Fail-as-is valve",
            "part_type": "valve",
        },
    ).json()

    changed = []
    for item in LOX_ITEMS:
        entry = dict(item)
        if entry["item_id"] == "fv201":
            entry["part_id"] = part["id"]
        if entry["item_id"] == "cv203":
            entry["tag"] = "CV-204"
        if entry["item_id"] == "trv201":
            continue  # deleted from the sheet
        changed.append(entry)
    client.put(f"/sheets/{sheet_id}", json={"index": {"items": changed, "lines": []}})

    stale = client.get(f"/fmea/{worksheet['id']}/stale").json()
    reasons = {(row["item_tag"] or row["item_id"], row["stale_reason"]) for row in stale}
    assert ("FV-201", "part_changed") in reasons
    assert ("CV-204", "retagged") in reasons
    assert ("trv201", "deleted") in reasons or ("TRV-201", "deleted") in reasons
    fv = next(row for row in stale if row["item_tag"] == "FV-201")
    assert "none → AMPH-VL-022" in fv["stale_detail"]
    assert fv["part_number"] == "AMPH-VL-022"
    view = client.get(f"/fmea/{worksheet['id']}").json()
    assert view["stale_count"] == len(stale)

    confirmed = client.post(f"/fmea/rows/{fv['id']}/confirm").json()
    assert confirmed["stale_reason"] is None
    client.post(f"/fmea/{worksheet['id']}/confirm-all")
    assert client.get(f"/fmea/{worksheet['id']}").json()["stale_count"] == 0
    # A revert is seen as a change again (the part goes back to none).
    client.put(f"/sheets/{sheet_id}", json={"index": {"items": LOX_ITEMS, "lines": []}})
    assert client.get(f"/fmea/{worksheet['id']}").json()["stale_count"] >= 1


def test_release_gate_release_diff_and_exports(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, sheet_id = _drawing_with_items(client, project_id, LOX_ITEMS[:1] + LOX_ITEMS[4:5])
    client.post(f"/drawings/{drawing_id}/revisions", json={"label": "B", "description": "Issued"})
    worksheet = _worksheet(client, project_id, drawing_id)
    assert worksheet["drawing_revision_label"] == "B"
    client.post(f"/fmea/{worksheet['id']}/generate", json={"categories": ["valve"]})
    rows = client.get(f"/fmea/{worksheet['id']}/rows").json()

    gate = client.get(f"/fmea/{worksheet['id']}/gate").json()
    assert gate["ready"] is False
    reasons = {blocker["reason"] for blocker in gate["blockers"]}
    assert any("needs a hazard" in reason for reason in reasons)
    blocked = client.post(f"/fmea/{worksheet['id']}/release", json={})
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["blockers"]

    hazard = client.post(
        f"/projects/{project_id}/hazards",
        json={
            "title": "Fast-fill overpressure",
            "severity_initial": "I",
            "likelihood_initial": "C",
        },
    ).json()
    requirement = client.post(
        "/requirements",
        json={
            "project_id": project_id,
            "key": "REQ-SAF-022",
            "title": "Interlock",
            "text": "Shall.",
            "requirement_type": "safety",
        },
    ).json()
    patches = []
    for row in rows:
        patch = {"id": row["id"], "occurrence": 3, "detection": 2}
        if row["severity"] is not None and row["severity"] >= 8:
            patch["hazard_id"] = hazard["id"]
        if row["detection_kind"] == "none":
            patch["detection_reason"] = "covered by relief"
        patches.append(patch)
    bulk = client.post(f"/fmea/{worksheet['id']}/rows/bulk", json={"rows": patches})
    assert bulk.status_code == 200, bulk.text
    assert all(row["rpn"] == row["severity"] * 6 for row in bulk.json())
    first = bulk.json()[0]
    linked = client.post(
        f"/fmea/rows/{first['id']}/controls", json={"type": "requirement", "id": requirement["id"]}
    )
    assert linked.status_code == 201, linked.text
    assert [control["label"] for control in linked.json()["controls"]] == ["REQ-SAF-022"]
    assert client.get(f"/hazards/{hazard['id']}").json()["causes"] >= 1

    assert client.get(f"/fmea/{worksheet['id']}/gate").json()["ready"] is True
    released = client.post(f"/fmea/{worksheet['id']}/release", json={"note": "CDR baseline"})
    assert released.status_code == 201, released.text
    assert released.json()["revision"] == 1
    assert released.json()["drawing_revision_label"] == "B"
    assert released.json()["row_count"] == len(rows)
    view = client.get(f"/fmea/{worksheet['id']}").json()
    assert view["status"] == "released" and view["revision"] == 1

    # Editing after release reopens the draft; the diff shows the change.
    edited = client.put(f"/fmea/rows/{first['id']}", json={"occurrence": 5, "cause": "seal wear"})
    assert edited.status_code == 200, edited.text
    assert client.get(f"/fmea/{worksheet['id']}").json()["status"] == "draft"
    diff = client.get(f"/fmea/{worksheet['id']}/diff").json()
    assert diff["against_revision"] == 1
    assert diff["added"] == [] and diff["removed"] == []
    assert len(diff["changed"]) == 1
    assert diff["changed"][0]["fields"]["occurrence"] == {"from": 3, "to": 5}
    assert "rpn" in diff["changed"][0]["fields"]
    client.post(f"/drawings/{drawing_id}/revisions", json={"label": "C", "description": "Reissued"})
    assert client.get(f"/fmea/{worksheet['id']}").json()["revision_drift"] is True

    xlsx = client.get(f"/fmea/{worksheet['id']}/export", params={"format": "xlsx"})
    assert xlsx.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx.content)).worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    header = next(row for row in values if row and row[0] == "Item")
    assert header[3] == "Failure mode" and "RPN" in header
    pdf = client.get(f"/fmea/{worksheet['id']}/export", params={"format": "pdf"})
    assert pdf.status_code == 200
    assert len(PdfReader(io.BytesIO(pdf.content)).pages) >= 1
    assert (
        client.get(f"/fmea/{worksheet['id']}/export", params={"format": "csv"}).status_code == 200
    )

    comment = client.post(
        f"/fmea/rows/{first['id']}/comments", json={"body": "Check the fail-safe position."}
    )
    assert comment.status_code == 201
    row_after = next(
        row
        for row in client.get(f"/fmea/{worksheet['id']}/rows").json()
        if row["id"] == first["id"]
    )
    assert (row_after["comment_count"], row_after["open_comment_count"]) == (1, 1)
    client.put(f"/fmea/comments/{comment.json()['id']}", json={"resolved": True})
    row_after = next(
        row
        for row in client.get(f"/fmea/{worksheet['id']}/rows").json()
        if row["id"] == first["id"]
    )
    assert row_after["open_comment_count"] == 0

    removed = client.delete(f"/fmea/rows/{first['id']}")
    assert removed.status_code == 204
    assert client.get(f"/fmea/{worksheet['id']}/diff").json()["removed"]
    gone = client.delete(f"/fmea/{worksheet['id']}")
    assert gone.status_code == 204
    assert client.get(f"/fmea/{worksheet['id']}").status_code == 404


def test_row_validation_and_manual_rows(client: TestClient) -> None:
    project_id = _project(client)
    drawing_id, sheet_id = _drawing_with_items(client, project_id, LOX_ITEMS[:2])
    worksheet = _worksheet(client, project_id, drawing_id)
    unknown = client.post(
        f"/fmea/{worksheet['id']}/rows",
        json={"sheet_id": sheet_id, "item_id": "nope", "failure_mode_text": "leaks"},
    )
    assert unknown.status_code == 422
    assert "not on sheet" in unknown.json()["detail"]
    empty = client.post(f"/fmea/{worksheet['id']}/rows", json={"cause": "x"})
    assert empty.status_code == 422
    out_of_range = client.post(
        f"/fmea/{worksheet['id']}/rows",
        json={
            "sheet_id": sheet_id,
            "item_id": "fv201",
            "failure_mode_text": "leaks",
            "severity": 11,
        },
    )
    assert out_of_range.status_code == 422
    manual = client.post(
        f"/fmea/{worksheet['id']}/rows",
        json={
            "subject_text": "Hold procedure",
            "failure_mode_text": "Operator skips drainback step",
            "severity": 9,
            "occurrence": 2,
            "detection": 4,
            "detection_kind": "procedure",
        },
    )
    assert manual.status_code == 201, manual.text
    assert manual.json()["rpn"] == 72 and manual.json()["item_exists"] is True
    bound = client.post(
        f"/fmea/{worksheet['id']}/rows",
        json={
            "sheet_id": sheet_id,
            "item_id": "fv201",
            "failure_mode_text": "Ice lock",
            "detected_by_item_id": "cv203",
        },
    )
    assert bound.status_code == 201, bound.text
    assert bound.json()["item_tag"] == "FV-201" and bound.json()["detected_by_tag"] == "CV-203"
    assert bound.json()["detection_kind"] == "instrument"
    viewer_blocked = client.post(
        f"/fmea/{worksheet['id']}/rows/bulk",
        json={"rows": [{"id": bound.json()["id"], "hazard_id": "missing"}]},
    )
    assert viewer_blocked.status_code == 422
