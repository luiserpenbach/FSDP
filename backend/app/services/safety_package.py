"""Safety review packages: one PDF and one XLSX per generation covering the
hazard log, risk matrix, FMEA rows by RPN, the verification matrix, open
actions, design rule findings and waivers, and the change log since the
previous package. Rendered server-side (SVG pages via cairosvg, openpyxl
workbook) so packages generate without a browser."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.config import settings as app_settings
from app.models import (
    Analysis,
    ChangeEvent,
    Drawing,
    DrawingSheet,
    FluidSystem,
    FmeaRelease,
    FmeaWorksheet,
    Hazard,
    Project,
    Requirement,
    SafetyPackage,
)
from app.services.drc import project_findings, verification_matrix
from app.services.export import svg_to_pdf
from app.services.fmea import export_rows, latest_revision_label, row_views
from app.services.hazards import matrix_counts, project_hazards
from app.services.lists import _cell
from app.services.safety_settings import get_settings
from app.services.svg_tables import key_value_page, table_pages

HAZARD_COLUMNS: list[tuple[str, str, float]] = [
    ("key", "Hazard", 0.6),
    ("title", "Title", 1.8),
    ("category", "Category", 0.8),
    ("operating_modes_text", "Modes", 0.9),
    ("risk_initial_text", "Initial", 0.6),
    ("risk_residual_text", "Residual", 0.6),
    ("controls_text", "Controls", 2.0),
    ("independent_controls", "Indep.", 0.45),
    ("fault_tolerance_required", "Req.", 0.45),
    ("computed_status", "Status", 0.7),
    ("owner", "Owner", 0.8),
]

FMEA_COLUMNS: list[tuple[str, str, float]] = [
    ("worksheet", "Worksheet", 0.9),
    ("item_tag", "Item", 0.8),
    ("failure_mode_title", "Failure mode", 1.4),
    ("cause", "Cause", 1.4),
    ("end_effect", "End effect", 1.6),
    ("detection_text", "Detected by", 0.9),
    ("severity", "S", 0.3),
    ("occurrence", "O", 0.3),
    ("detection", "D", 0.3),
    ("rpn", "RPN", 0.45),
    ("controls_text", "Controls", 1.1),
    ("hazard_key", "Hazard", 0.6),
    ("recommended_action", "Action", 1.2),
    ("action_status", "Action status", 0.7),
]

REQUIREMENT_COLUMNS: list[tuple[str, str, float]] = [
    ("key", "Requirement", 0.8),
    ("title", "Title", 2.2),
    ("category", "Category", 0.7),
    ("verification_method", "Method", 0.8),
    ("verification_status", "Verification", 0.8),
    ("verdict", "Design rule", 0.7),
    ("checked", "Checked", 0.5),
    ("failed", "Failed", 0.5),
    ("evidence_text", "Evidence", 1.1),
    ("hazard_keys_text", "Hazards", 0.8),
    ("safety_critical", "Safety critical", 0.6),
]

ACTION_COLUMNS: list[tuple[str, str, float]] = [
    ("kind", "Kind", 0.6),
    ("ref", "Ref", 0.8),
    ("title", "Title", 2.4),
    ("detail", "Open item", 2.4),
    ("owner", "Owner", 0.9),
    ("status", "Status", 0.8),
]

DRC_COLUMNS: list[tuple[str, str, float]] = [
    ("drawing_number", "Drawing", 0.8),
    ("sheet_no", "Sheet", 0.4),
    ("zone", "Zone", 0.4),
    ("rule", "Rule", 1.0),
    ("severity", "Severity", 0.6),
    ("message", "Finding", 3.0),
    ("hazard_key", "Hazard", 0.6),
    ("waived_text", "Waived", 0.5),
    ("waiver_reason", "Waiver reason", 1.4),
]

CHANGE_COLUMNS: list[tuple[str, str, float]] = [
    ("when", "When", 0.9),
    ("object_type", "Object", 0.8),
    ("action", "Action", 0.6),
    ("summary", "Summary", 3.6),
    ("actor", "Actor", 1.0),
]

XLSX_SHEETS: list[tuple[str, list[tuple[str, str, float]], str]] = [
    ("Hazards", HAZARD_COLUMNS, "hazards"),
    ("FMEA", FMEA_COLUMNS, "fmea_rows"),
    ("Requirements", REQUIREMENT_COLUMNS, "requirements"),
    ("Open actions", ACTION_COLUMNS, "actions"),
    ("Design rules", DRC_COLUMNS, "drc"),
    ("Change log", CHANGE_COLUMNS, "changes"),
]


def safety_files_root() -> Path:
    path = Path(app_settings.safety_files_dir)
    if not path.is_absolute():
        path = Path.cwd() / path
    path.mkdir(parents=True, exist_ok=True)
    return path


def package_file(package: SafetyPackage, kind: str) -> Path | None:
    relative = package.pdf_path if kind == "pdf" else package.xlsx_path
    if not relative:
        return None
    return safety_files_root() / relative


def _stamp(value: datetime | None) -> str:
    if value is None:
        return "—"
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC")


def resolve_scope(
    db: Session,
    project: Project,
    *,
    drawing_ids: list[str] | None,
    worksheet_ids: list[str] | None,
    system_ids: list[str] | None,
) -> dict[str, Any]:
    """The systems, drawings (with current revision), and worksheets (with
    revision and status) the package covers. Empty selections mean everything."""
    systems = list(
        db.scalars(
            select(FluidSystem)
            .where(FluidSystem.project_id == project.id)
            .order_by(FluidSystem.name)
        )
    )
    if system_ids:
        systems = [system for system in systems if system.id in set(system_ids)]
    drawings = list(
        db.scalars(
            select(Drawing)
            .where(Drawing.project_id == project.id)
            .options(selectinload(Drawing.revisions), selectinload(Drawing.sheets))
            .order_by(Drawing.number)
        )
    )
    if drawing_ids:
        drawings = [drawing for drawing in drawings if drawing.id in set(drawing_ids)]
    worksheets = list(
        db.scalars(
            select(FmeaWorksheet)
            .where(FmeaWorksheet.project_id == project.id)
            .options(selectinload(FmeaWorksheet.rows), selectinload(FmeaWorksheet.releases))
            .order_by(FmeaWorksheet.title)
        )
    )
    if worksheet_ids:
        worksheets = [sheet for sheet in worksheets if sheet.id in set(worksheet_ids)]
    return {
        "systems": [{"id": system.id, "name": system.name} for system in systems],
        "drawings": [
            {
                "id": drawing.id,
                "number": drawing.number,
                "title": drawing.title,
                "revision": latest_revision_label(drawing),
                "sheets": len(drawing.sheets),
            }
            for drawing in drawings
        ],
        "worksheets": [
            {
                "id": sheet.id,
                "title": sheet.title,
                "revision": sheet.revision,
                "status": sheet.status,
                "drawing_revision": sheet.drawing_revision_label,
                "rows": len(sheet.rows),
            }
            for sheet in worksheets
        ],
        "_drawings": drawings,
        "_worksheets": worksheets,
    }


def _control_status(control: dict[str, Any]) -> str:
    if control.get("verification_status"):
        return str(control["verification_status"])
    return "covered" if control.get("covered") else "uncovered"


def _hazard_rows(views: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for view in views:
        rows.append(
            {
                **view,
                "operating_modes_text": ", ".join(view.get("operating_modes") or []),
                "risk_initial_text": (
                    f"{view.get('severity_initial') or '-'}{view.get('likelihood_initial') or '-'}"
                    f" ({view.get('risk_initial') or 'unrated'})"
                ),
                "risk_residual_text": (
                    f"{view.get('severity_residual') or '-'}"
                    f"{view.get('likelihood_residual') or '-'}"
                    f" ({view.get('risk_residual') or 'unrated'})"
                ),
                "controls_text": "; ".join(
                    f"{control['label']} [{_control_status(control)}]"
                    for control in view.get("controls", [])
                ),
            }
        )
    return rows


def _requirement_rows(matrix: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in matrix["rows"]:
        evidence = row.get("evidence") or {}
        rows.append(
            {
                **row,
                "evidence_text": ", ".join(f"{kind} ×{count}" for kind, count in evidence.items())
                or "none",
                "hazard_keys_text": ", ".join(row.get("hazards") or []),
                "safety_critical": bool(row.get("safety_critical")),
            }
        )
    return rows


def _fmea_rows(db: Session, worksheets: list[FmeaWorksheet]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for worksheet in worksheets:
        for row in export_rows(row_views(db, worksheet, list(worksheet.rows))):
            rows.append({**row, "worksheet": worksheet.title})
    rows.sort(key=lambda row: -(row.get("rpn") or 0))
    return rows


def _actions(
    hazards: list[dict[str, Any]],
    fmea_rows: list[dict[str, Any]],
    requirements: list[dict[str, Any]],
    findings: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    actions: list[dict[str, Any]] = []
    for hazard in hazards:
        if hazard["computed_status"] in {"accepted", "closed"}:
            continue
        if hazard.get("severity_initial") in {"I", "II"} or hazard["computed_status"] == "open":
            actions.append(
                {
                    "kind": "hazard",
                    "ref": hazard["key"],
                    "title": hazard["title"],
                    "detail": (
                        f"{hazard['computed_status']}: {hazard['controls_verified']} of "
                        f"{hazard['controls_total']} controls verified, "
                        f"{hazard['independent_controls']} independent of "
                        f"{hazard['fault_tolerance_required']} required"
                    ),
                    "owner": hazard.get("owner"),
                    "status": hazard["computed_status"],
                }
            )
    for row in fmea_rows:
        if row.get("recommended_action") and row.get("action_status") not in {"closed", "done"}:
            actions.append(
                {
                    "kind": "fmea",
                    "ref": row.get("item_tag"),
                    "title": row.get("failure_mode_title") or row.get("failure_mode_text") or "",
                    "detail": row.get("recommended_action"),
                    "owner": row.get("action_owner"),
                    "status": row.get("action_status") or "open",
                }
            )
    for requirement in requirements:
        if requirement.get("safety_critical") and requirement.get("verification_status") not in {
            "verified",
            "waived",
        }:
            actions.append(
                {
                    "kind": "requirement",
                    "ref": requirement["key"],
                    "title": requirement["title"],
                    "detail": f"safety-critical, verification {requirement['verification_status']}",
                    "owner": requirement.get("owner"),
                    "status": requirement["verification_status"],
                }
            )
    for finding in findings:
        if finding["severity"] == "error" and not finding["waived"]:
            actions.append(
                {
                    "kind": "design_rule",
                    "ref": f"{finding['drawing_number']} sh {finding['sheet_no']}",
                    "title": finding["rule"].replace("_", " "),
                    "detail": finding["message"],
                    "owner": None,
                    "status": "open",
                }
            )
    return actions


def _project_object_ids(db: Session, project: Project, scope: dict[str, Any]) -> set[str]:
    ids: set[str] = {project.id}
    for model in (Hazard, Requirement, FmeaWorksheet, Analysis, Drawing):
        ids.update(
            db.scalars(select(model.id).where(model.project_id == project.id))  # type: ignore[attr-defined]
        )
    worksheet_ids = [sheet.id for sheet in scope["_worksheets"]]
    if worksheet_ids:
        ids.update(
            db.scalars(select(FmeaRelease.id).where(FmeaRelease.worksheet_id.in_(worksheet_ids)))
        )
    drawing_ids = [drawing.id for drawing in scope["_drawings"]]
    if drawing_ids:
        ids.update(
            db.scalars(select(DrawingSheet.id).where(DrawingSheet.drawing_id.in_(drawing_ids)))
        )
    return ids


def _naive_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)


def change_log(
    db: Session, project: Project, scope: dict[str, Any], since: datetime | None, limit: int = 300
) -> list[dict[str, Any]]:
    """Change events on the project's safety objects since ``since`` (inclusive:
    server timestamps have second resolution). Filtered in Python so SQLite and
    PostgreSQL agree on the boundary."""
    ids = _project_object_ids(db, project, scope)
    boundary = _naive_utc(since) if since is not None else None
    rows = []
    for event in db.scalars(
        select(ChangeEvent).order_by(ChangeEvent.created_at.desc()).limit(limit * 4)
    ):
        if event.object_id not in ids or event.object_type == "safety_package":
            continue
        if boundary is not None and _naive_utc(event.created_at) < boundary:
            continue
        rows.append(
            {
                "when": _stamp(event.created_at),
                "object_type": event.object_type,
                "action": event.action,
                "summary": event.summary,
                "actor": event.actor,
            }
        )
        if len(rows) >= limit:
            break
    return rows


def _matrix_rows(matrix: dict[str, Any], key: str) -> list[dict[str, Any]]:
    likelihoods = [entry["code"] for entry in matrix["likelihood_scale"]]
    rows = []
    for severity in matrix["severity_scale"]:
        counts = matrix[key].get(severity["code"], {})
        rows.append(
            {
                "severity": f"{severity['code']} {severity['name']}",
                **{code: counts.get(code, 0) for code in likelihoods},
            }
        )
    return rows


def build_pages(project: Project, package: SafetyPackage, data: dict[str, Any]) -> list[str]:
    scope = data["scope"]
    summary = data["summary"]
    subtitle = (
        f"{project.name} · generated {_stamp(package.generated_at)} by {package.generated_by}"
    )
    cover_sections: list[tuple[str, list[tuple[str, str]]]] = [
        (
            "Scope",
            [
                ("Systems", ", ".join(s["name"] for s in scope["systems"]) or "all"),
                (
                    "Drawings",
                    "; ".join(
                        f"{d['number']} rev {d['revision'] or '-'} ({d['sheets']} sheets)"
                        for d in scope["drawings"]
                    )
                    or "none",
                ),
                (
                    "Worksheets",
                    "; ".join(
                        f"{w['title']} rev {w['revision']} ({w['status']}, drawing rev "
                        f"{w['drawing_revision'] or '-'})"
                        for w in scope["worksheets"]
                    )
                    or "none",
                ),
                ("Change log since", _stamp(package.change_log_from)),
            ],
        ),
        (
            "Summary",
            [
                (
                    "Hazards",
                    f"{summary['hazards']} ({summary['hazards_open']} open, "
                    f"{summary['hazards_controlled']} controlled, "
                    f"{summary['hazards_accepted']} accepted, {summary['hazards_closed']} closed)",
                ),
                ("Severity I–II not accepted", str(summary["hazards_high_open"])),
                (
                    "Safety requirements verified",
                    f"{summary['requirements_verified']} of {summary['requirements_safety']}",
                ),
                (
                    "FMEA rows",
                    f"{summary['fmea_rows']} ({summary['fmea_over_threshold']} at or above RPN "
                    f"{summary['rpn_threshold']}, {summary['fmea_stale']} stale)",
                ),
                (
                    "Design rule findings",
                    f"{summary['drc_errors']} errors, {summary['drc_warnings']} warnings, "
                    f"{summary['drc_waived']} waived",
                ),
                ("Open actions", str(summary["actions"])),
                ("Changes since previous package", str(summary["changes"])),
            ],
        ),
    ]
    pages = [key_value_page(package.title, subtitle, cover_sections)]
    matrix = data["matrix"]
    likelihoods = [entry["code"] for entry in matrix["likelihood_scale"]]
    matrix_columns = [("severity", "Severity", 1.6)] + [(code, code, 0.5) for code in likelihoods]
    pages += table_pages(
        "Risk matrix · initial",
        f"{subtitle} · hazards by initial severity × likelihood",
        matrix_columns,
        _matrix_rows(matrix, "initial"),
    )
    pages += table_pages(
        "Risk matrix · residual",
        f"{subtitle} · hazards by residual severity × likelihood",
        matrix_columns,
        _matrix_rows(matrix, "residual"),
    )
    pages += table_pages("Hazard log", subtitle, HAZARD_COLUMNS, data["hazards"])
    pages += table_pages(
        "FMEA rows by RPN", subtitle, FMEA_COLUMNS, data["fmea_rows"], footer="Sorted by RPN."
    )
    pages += table_pages("Verification matrix", subtitle, REQUIREMENT_COLUMNS, data["requirements"])
    pages += table_pages("Open actions", subtitle, ACTION_COLUMNS, data["actions"])
    pages += table_pages("Design rule findings and waivers", subtitle, DRC_COLUMNS, data["drc"])
    pages += table_pages(
        f"Change log since {_stamp(package.change_log_from)}",
        subtitle,
        CHANGE_COLUMNS,
        data["changes"],
    )
    return pages


def build_xlsx(project: Project, package: SafetyPackage, data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDE3EA")
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Package", package.title])
    summary_sheet.append(["Project", project.name])
    summary_sheet.append(["Generated", _stamp(package.generated_at)])
    summary_sheet.append(["Generated by", package.generated_by or ""])
    summary_sheet.append(["Change log since", _stamp(package.change_log_from)])
    summary_sheet.append([])
    for key, value in data["summary"].items():
        summary_sheet.append([key.replace("_", " ").capitalize(), value])
    for row in summary_sheet.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 60
    for title, columns, key in XLSX_SHEETS:
        sheet = workbook.create_sheet(title[:31])
        sheet.append([label for _, label, _ in columns])
        for index in range(1, len(columns) + 1):
            cell = sheet.cell(row=1, column=index)
            cell.font = bold
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        rows = data[key]
        for row in rows:
            sheet.append(
                [
                    row.get(col) if isinstance(row.get(col), int | float) else _cell(row.get(col))
                    for col, _, _ in columns
                ]
            )
        for index, (col, label, _) in enumerate(columns, start=1):
            width = max([len(label)] + [len(_cell(row.get(col))) for row in rows] + [4])
            sheet.column_dimensions[get_column_letter(index)].width = min(60, width + 2)
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def collect(db: Session, project: Project, scope: dict[str, Any], since: datetime | None) -> dict:
    """Everything a package renders, as plain rows."""
    safety = get_settings(db, project)
    hazard_views = project_hazards(db, project)
    matrix = matrix_counts(db, project)
    requirement_rows = _requirement_rows(verification_matrix(db, project))
    fmea_rows = _fmea_rows(db, scope["_worksheets"])
    drc = project_findings(db, project)
    drawing_ids = {drawing["id"] for drawing in scope["drawings"]}
    findings = [
        {**finding, "waived_text": "yes" if finding["waived"] else ""}
        for finding in drc["findings"]
        if finding["drawing_id"] in drawing_ids
    ]
    actions = _actions(hazard_views, fmea_rows, requirement_rows, findings)
    changes = change_log(db, project, scope, since)
    threshold = int(safety.get("fmea_rpn_threshold") or 100)
    safety_requirements = [row for row in requirement_rows if row.get("safety_critical")]
    summary = {
        "hazards": len(hazard_views),
        "hazards_open": sum(1 for h in hazard_views if h["computed_status"] == "open"),
        "hazards_controlled": sum(1 for h in hazard_views if h["computed_status"] == "controlled"),
        "hazards_accepted": sum(1 for h in hazard_views if h["computed_status"] == "accepted"),
        "hazards_closed": sum(1 for h in hazard_views if h["computed_status"] == "closed"),
        "hazards_high_open": sum(
            1
            for h in hazard_views
            if h.get("severity_initial") in {"I", "II"}
            and h["computed_status"] not in {"accepted", "closed"}
        ),
        "requirements_safety": len(safety_requirements),
        "requirements_verified": sum(
            1 for row in safety_requirements if row.get("verification_status") == "verified"
        ),
        "fmea_rows": len(fmea_rows),
        "fmea_over_threshold": sum(1 for row in fmea_rows if (row.get("rpn") or 0) >= threshold),
        "fmea_stale": sum(1 for row in fmea_rows if row.get("stale_reason")),
        "rpn_threshold": threshold,
        "drc_errors": sum(1 for f in findings if f["severity"] == "error" and not f["waived"]),
        "drc_warnings": sum(1 for f in findings if f["severity"] == "warning" and not f["waived"]),
        "drc_waived": sum(1 for f in findings if f["waived"]),
        "actions": len(actions),
        "changes": len(changes),
    }
    return {
        "scope": {key: value for key, value in scope.items() if not key.startswith("_")},
        "summary": summary,
        "matrix": matrix,
        "hazards": _hazard_rows(hazard_views),
        "fmea_rows": fmea_rows,
        "requirements": requirement_rows,
        "actions": actions,
        "drc": findings,
        "changes": changes,
    }


def generate(
    db: Session,
    project: Project,
    *,
    title: str | None,
    drawing_ids: list[str] | None,
    worksheet_ids: list[str] | None,
    system_ids: list[str] | None,
    actor: str | None,
) -> SafetyPackage:
    """Build, render, store, and record a package. Files land under the
    configured safety files directory; the row keeps their relative paths."""
    previous = db.scalar(
        select(SafetyPackage)
        .where(SafetyPackage.project_id == project.id)
        .order_by(SafetyPackage.generated_at.desc())
    )
    scope = resolve_scope(
        db, project, drawing_ids=drawing_ids, worksheet_ids=worksheet_ids, system_ids=system_ids
    )
    now = datetime.now(UTC)
    count = len(
        list(db.scalars(select(SafetyPackage.id).where(SafetyPackage.project_id == project.id)))
    )
    package = SafetyPackage(
        project_id=project.id,
        title=title or f"Safety review package {count + 1} · {project.name}",
        generated_by=actor,
        generated_at=now,
        change_log_from=previous.created_at if previous else None,
    )
    data = collect(db, project, scope, package.change_log_from)
    package.scope = data["scope"]
    package.summary = data["summary"]
    db.add(package)
    db.flush()
    pages = build_pages(project, package, data)
    pdf = svg_to_pdf(pages[0], pages[1:])
    xlsx = build_xlsx(project, package, data)
    folder = safety_files_root() / project.id
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{package.id}.pdf").write_bytes(pdf)
    (folder / f"{package.id}.xlsx").write_bytes(xlsx)
    package.pdf_path = f"{project.id}/{package.id}.pdf"
    package.xlsx_path = f"{project.id}/{package.id}.xlsx"
    return package


def package_view(package: SafetyPackage) -> dict[str, Any]:
    return {
        "id": package.id,
        "project_id": package.project_id,
        "title": package.title,
        "scope": package.scope,
        "summary": package.summary,
        "generated_by": package.generated_by,
        "generated_at": package.generated_at,
        "change_log_from": package.change_log_from,
        "pdf_url": f"/safety/packages/{package.id}/pdf",
        "xlsx_url": f"/safety/packages/{package.id}/xlsx",
        "created_at": package.created_at,
    }
