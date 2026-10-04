"""Bulk import of parts and requirements from CSV, XLSX or pasted spreadsheet rows.

This module reads a table (an uploaded file or JSON rows sent after a paste from
Excel), maps its column headers onto model fields through forgiving aliases, and
plans the import row by row: every row is validated with the same Pydantic models
the single-object endpoints use (PartCreate/PartUpdate, RequirementCreate/
RequirementUpdate), so the import cannot drift from the API's rules. The plan is a
dry-run report; the API layer applies it (all or nothing) when asked to commit.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Part, Requirement
from app.schemas import PartCreate, PartUpdate, RequirementCreate, RequirementUpdate

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
MAX_IMPORT_COLUMNS = 100
# Rows scanned for the header row: exports often start with a title block.
HEADER_SCAN_ROWS = 20

ImportMode = Literal["create_only", "upsert"]
FieldKind = Literal["text", "float", "bool", "enum", "json"]

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}


class ImportFormatError(ValueError):
    """The input cannot be read as a table (bad encoding, not a workbook, too many rows)."""


# --------------------------------------------------------------------------- fields


@dataclass(frozen=True)
class FieldSpec:
    name: str
    aliases: tuple[str, ...] = ()
    kind: FieldKind = "text"
    example: Any = None
    in_template: bool = True


@dataclass(frozen=True)
class EntitySpec:
    entity: str
    label: str
    model: type
    key_field: str
    create_schema: type[BaseModel]
    update_schema: type[BaseModel]
    fields: tuple[FieldSpec, ...]
    # Fields bulk edit may set on many rows at once (identity and prose excluded).
    bulk_fields: frozenset[str]
    # Fields that may be cleared (set to null) by bulk edit.
    nullable_fields: frozenset[str]

    def field(self, name: str) -> FieldSpec:
        return next(spec for spec in self.fields if spec.name == name)

    @property
    def alias_lookup(self) -> dict[str, str]:
        lookup: dict[str, str] = {}
        for spec in self.fields:
            for alias in (spec.name, *spec.aliases):
                lookup.setdefault(normalize_header(alias), spec.name)
        return lookup

    @property
    def key_aliases(self) -> set[str]:
        lookup = self.alias_lookup
        return {alias for alias, name in lookup.items() if name == self.key_field}

    @property
    def template_fields(self) -> list[FieldSpec]:
        return [spec for spec in self.fields if spec.in_template]


def normalize_header(value: Any) -> str:
    """Case, space, underscore and punctuation insensitive header key ("Part #" -> "partno")."""
    text = str(value or "").casefold().replace("#", "no")
    return re.sub(r"[^0-9a-z]+", "", text)


PART_SPEC = EntitySpec(
    entity="part",
    label="parts",
    model=Part,
    key_field="part_number",
    create_schema=PartCreate,
    update_schema=PartUpdate,
    fields=(
        FieldSpec(
            "part_number",
            (
                "Part Number",
                "Part No",
                "Part #",
                "PN",
                "P/N",
                "Part",
                "Part Name",
                "Name",
                "Item Number",
            ),
            example="PV-1001",
        ),
        FieldSpec("revision", ("Rev", "Revision", "Part Revision"), example="A"),
        FieldSpec("description", ("Description", "Desc"), example="Ball valve 1/2 in, 316L"),
        FieldSpec(
            "manufacturer",
            ("Mfr", "Mfg", "Manufacturer", "Maker", "Vendor", "Supplier"),
            example="Acme Valves",
        ),
        FieldSpec("part_type", ("Type", "Part Type", "Category", "Kind"), example="valve"),
        FieldSpec("source_type", ("Source", "Source Type"), kind="enum", example="vendor"),
        FieldSpec("material", ("Material", "Mat", "Body Material"), example="316L"),
        FieldSpec(
            "pressure_rating_bar",
            (
                "Pressure Rating (bar)",
                "Pressure Rating",
                "Pressure",
                "Rating (bar)",
                "Bar",
                "MAWP",
                "MAWP (bar)",
                "Max Pressure (bar)",
            ),
            kind="float",
            example=206,
        ),
        FieldSpec(
            "temperature_min_c",
            ("Temp Min", "Min Temp", "Min Temperature", "Temperature Min (C)", "Tmin", "T min"),
            kind="float",
            example=-40,
        ),
        FieldSpec(
            "temperature_max_c",
            ("Temp Max", "Max Temp", "Max Temperature", "Temperature Max (C)", "Tmax", "T max"),
            kind="float",
            example=120,
        ),
        FieldSpec("cv", ("Cv", "Flow Coefficient"), kind="float", example=1.2),
        FieldSpec(
            "mass_kg", ("Mass", "Mass (kg)", "Weight", "Weight (kg)"), kind="float", example=0.45
        ),
        FieldSpec(
            "certification_status",
            ("Cert", "Certification", "Certification Status", "Cert Status"),
            kind="enum",
            example="unreviewed",
        ),
        FieldSpec(
            "qualification_status",
            ("Qual", "Qualification", "Qualification Status", "Qual Status"),
            kind="enum",
            example="unqualified",
        ),
        FieldSpec(
            "lifecycle_status",
            ("Lifecycle", "Lifecycle Status", "Status", "Life Cycle"),
            kind="enum",
            example="draft",
        ),
        FieldSpec("preferred", ("Preferred", "Preferred Part"), kind="bool", example="no"),
        FieldSpec("notes", ("Notes", "Note", "Comments", "Comment", "Remarks"), example=""),
        FieldSpec("dimensions", ("Dimensions",), kind="json", in_template=False),
        FieldSpec("metadata", ("Metadata",), kind="json", in_template=False),
    ),
    bulk_fields=frozenset(
        {
            "manufacturer",
            "part_type",
            "source_type",
            "material",
            "pressure_rating_bar",
            "temperature_min_c",
            "temperature_max_c",
            "cv",
            "mass_kg",
            "certification_status",
            "qualification_status",
            "lifecycle_status",
            "preferred",
            "notes",
        }
    ),
    nullable_fields=frozenset(
        {
            "manufacturer",
            "material",
            "pressure_rating_bar",
            "temperature_min_c",
            "temperature_max_c",
            "cv",
            "mass_kg",
            "notes",
        }
    ),
)

REQUIREMENT_SPEC = EntitySpec(
    entity="requirement",
    label="requirements",
    model=Requirement,
    key_field="key",
    create_schema=RequirementCreate,
    update_schema=RequirementUpdate,
    fields=(
        FieldSpec(
            "key",
            ("ID", "Key", "Req ID", "Requirement ID", "Req", "Req No", "Requirement Key", "Ref"),
            example="REQ-001",
        ),
        FieldSpec("title", ("Title", "Name", "Summary", "Heading"), example="Wetted materials"),
        FieldSpec(
            "text",
            ("Text", "Description", "Requirement", "Requirement Text", "Statement", "Shall"),
            example="All wetted parts shall be 316L stainless steel.",
        ),
        FieldSpec(
            "requirement_type",
            ("Type", "Requirement Type", "Req Type", "Category"),
            example="materials",
        ),
        FieldSpec(
            "verification_method",
            ("Verification Method", "Verification", "Verify", "Method", "VM", "V Method"),
            example="inspection",
        ),
        FieldSpec("status", ("Status", "State"), example="draft"),
        FieldSpec("owner", ("Owner", "Responsible", "Assignee"), example=""),
        FieldSpec(
            "constraint",
            ("Constraint", "Check", "Rule"),
            kind="json",
            example='{"kind": "material_in", "values": ["316L"]}',
        ),
    ),
    bulk_fields=frozenset(
        {"requirement_type", "verification_method", "status", "owner", "constraint"}
    ),
    nullable_fields=frozenset({"verification_method", "owner", "constraint"}),
)

ENTITY_SPECS = {"part": PART_SPEC, "requirement": REQUIREMENT_SPEC}


# --------------------------------------------------------------------------- reading


@dataclass
class Table:
    headers: list[str]
    # (1-based data row number, cells); blank rows keep their number but are dropped.
    rows: list[tuple[int, list[Any]]]
    decimal_comma: bool = False


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        # Excel's "Unicode Text" export: UTF-16 with BOM, tab separated.
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _detect_delimiter(text: str) -> str:
    sample = text[:16384]
    first_line = sample.splitlines()[0] if sample.splitlines() else ""
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        if delimiter in first_line:
            return delimiter
    except csv.Error:
        pass
    counts = {candidate: first_line.count(candidate) for candidate in (",", ";", "\t", "|")}
    best = max(counts, key=lambda candidate: counts[candidate])
    return best if counts[best] else ","


def _row_limit_error() -> ImportFormatError:
    return ImportFormatError(f"Too many rows: the limit is {MAX_IMPORT_ROWS} data rows per import")


def read_csv_rows(content: bytes) -> tuple[list[list[Any]], str]:
    text = _decode_text(content)
    if not text.strip():
        return [], ","
    delimiter = _detect_delimiter(text)
    rows: list[list[Any]] = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            rows.append(row[:MAX_IMPORT_COLUMNS])
            if len(rows) > MAX_IMPORT_ROWS + HEADER_SCAN_ROWS:
                raise _row_limit_error()
    except csv.Error as exc:
        raise ImportFormatError(f"Could not read the CSV file: {exc}") from exc
    return rows, delimiter


def read_xlsx_rows(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many exception types for corrupt files
        raise ImportFormatError("The file is not a readable .xlsx workbook") from exc
    try:
        sheet = workbook.active or workbook.worksheets[0]
        rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row[:MAX_IMPORT_COLUMNS]))
            if len(rows) > MAX_IMPORT_ROWS + HEADER_SCAN_ROWS:
                raise _row_limit_error()
        return rows
    finally:
        workbook.close()


def _is_blank(cells: Iterable[Any]) -> bool:
    return all(cell_text(cell) is None for cell in cells)


def build_table(raw_rows: list[list[Any]], key_aliases: set[str]) -> Table:
    """Find the header row (the first row naming the key column, else the first non-blank)."""
    header_index: int | None = None
    first_non_blank: int | None = None
    for index, row in enumerate(raw_rows[: HEADER_SCAN_ROWS + 1]):
        if _is_blank(row):
            continue
        if first_non_blank is None:
            first_non_blank = index
        if any(normalize_header(cell) in key_aliases for cell in row if cell is not None):
            header_index = index
            break
    if header_index is None:
        header_index = first_non_blank
    if header_index is None:
        return Table(headers=[], rows=[])
    headers = [cell_text(cell) or "" for cell in raw_rows[header_index]]
    while headers and not headers[-1]:
        headers.pop()
    rows: list[tuple[int, list[Any]]] = []
    for number, row in enumerate(raw_rows[header_index + 1 :], start=1):
        if _is_blank(row):
            continue
        rows.append((number, list(row)))
    if len(rows) > MAX_IMPORT_ROWS:
        raise _row_limit_error()
    return Table(headers=headers, rows=rows)


def table_from_upload(content: bytes, filename: str, spec_key_aliases: set[str]) -> Table:
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if suffix in XLSX_SUFFIXES or (not suffix and content.startswith(b"PK")):
        return build_table(read_xlsx_rows(content), spec_key_aliases)
    if suffix == ".xls":
        raise ImportFormatError("Legacy .xls workbooks are not supported; save as .xlsx or .csv")
    if suffix in CSV_SUFFIXES or not suffix:
        raw_rows, delimiter = read_csv_rows(content)
        table = build_table(raw_rows, spec_key_aliases)
        # A semicolon-separated file comes from a locale that writes 2,5 for 2.5.
        table.decimal_comma = delimiter == ";"
        return table
    raise ImportFormatError("Unsupported file type; upload a .csv or .xlsx file")


def table_from_json(
    headers: list[str] | None, rows: list[Any], decimal_comma: bool = False
) -> Table:
    """Rows from a paste: lists of cells (with `headers` or a header first row) or dicts."""
    if len(rows) > MAX_IMPORT_ROWS + 1:
        raise _row_limit_error()
    if rows and all(isinstance(row, dict) for row in rows):
        keys: list[str] = list(headers or [])
        for row in rows:
            for key in row:
                if key not in keys:
                    keys.append(key)
        keys = keys[:MAX_IMPORT_COLUMNS]
        table_rows = [
            (number, [row.get(key) for key in keys])
            for number, row in enumerate(rows, start=1)
            if not _is_blank(row.values())
        ]
        return Table(
            headers=[str(key) for key in keys], rows=table_rows, decimal_comma=decimal_comma
        )
    if not all(isinstance(row, list) for row in rows):
        raise ImportFormatError("rows must be all lists of cells or all objects")
    if headers is None:
        if not rows:
            return Table(headers=[], rows=[], decimal_comma=decimal_comma)
        headers, rows = [cell_text(cell) or "" for cell in rows[0]], rows[1:]
    table_rows = [
        (number, list(row[:MAX_IMPORT_COLUMNS]))
        for number, row in enumerate(rows, start=1)
        if not _is_blank(row)
    ]
    if len(table_rows) > MAX_IMPORT_ROWS:
        raise _row_limit_error()
    return Table(
        headers=[str(header or "").strip() for header in headers[:MAX_IMPORT_COLUMNS]],
        rows=table_rows,
        decimal_comma=decimal_comma,
    )


# --------------------------------------------------------------------------- cells


def cell_text(value: Any) -> str | None:
    """A spreadsheet cell as trimmed text, None when blank."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, dict | list):
        return json.dumps(value)
    text = str(value).strip()
    # Undo the formula-injection guard our CSV/XLSX exports add ('=..., '-5, ...).
    if len(text) >= 2 and text[0] == "'" and text[1] in "=+-@":
        text = text[1:]
    return text or None


_THOUSANDS_COMMA = re.compile(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?")


def normalize_number(text: str, decimal_comma: bool) -> str:
    compact = text.replace(" ", "").replace(" ", "").replace(" ", "")
    if decimal_comma:
        if "," in compact:
            return compact.replace(".", "").replace(",", ".")
        return compact
    if "," in compact:
        if _THOUSANDS_COMMA.fullmatch(compact):
            return compact.replace(",", "")
        if "." not in compact and compact.count(",") == 1:
            return compact.replace(",", ".")
    return compact


_TRUE_MARKS = {"x", "✓", "✔"}


def normalize_enum(text: str) -> str:
    return re.sub(r"[\s\-]+", "_", text.strip().lower())


def coerce_cell(spec: FieldSpec, text: str, decimal_comma: bool) -> Any:
    """Turn cell text into the value fed to the Pydantic model (which does the validation)."""
    if spec.kind == "float":
        return normalize_number(text, decimal_comma)
    if spec.kind == "bool":
        return "true" if text.strip().lower() in _TRUE_MARKS else text
    if spec.kind == "enum":
        return normalize_enum(text)
    if spec.kind == "json":
        return json.loads(text)
    return text


# --------------------------------------------------------------------------- mapping


@dataclass
class Mapping:
    mapping: dict[str, str | None]
    columns: list[tuple[int, str]]  # (column index, field)
    warnings: list[str] = field(default_factory=list)


def resolve_mapping(spec: EntitySpec, headers: list[str]) -> Mapping:
    lookup = spec.alias_lookup
    mapping: dict[str, str | None] = {}
    columns: list[tuple[int, str]] = []
    claimed: dict[str, str] = {}
    warnings: list[str] = []
    for index, header in enumerate(headers):
        if not header:
            continue
        if header in mapping:
            warnings.append(f"Duplicate column '{header}' ignored")
            continue
        target = lookup.get(normalize_header(header))
        if target is not None and target in claimed:
            warnings.append(
                f"Column '{header}' also maps to {target}; using column '{claimed[target]}'"
            )
            target = None
        mapping[header] = target
        if target is not None:
            claimed[target] = header
            columns.append((index, target))
    return Mapping(mapping=mapping, columns=columns, warnings=warnings)


# --------------------------------------------------------------------------- planning


@dataclass
class RowPlan:
    row: int
    action: str
    key: str | None
    errors: list[dict[str, str]] = field(default_factory=list)
    changes: dict[str, list[Any]] = field(default_factory=dict)
    target: Any | None = None
    create_data: dict[str, Any] | None = None

    def report(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "action": self.action,
            "key": self.key,
            "id": getattr(self.target, "id", None),
            "errors": self.errors,
            "changes": self.changes,
        }


@dataclass
class ImportPlan:
    spec: EntitySpec
    mode: ImportMode
    mapping: Mapping
    rows: list[RowPlan]
    project_id: str | None = None

    @property
    def has_errors(self) -> bool:
        return any(row.action == "error" for row in self.rows)

    def summary(self) -> dict[str, int]:
        counts = {"create": 0, "update": 0, "unchanged": 0, "error": 0}
        for row in self.rows:
            counts[row.action] += 1
        return counts

    def report(self, *, dry_run: bool, committed: bool) -> dict[str, Any]:
        return {
            "entity": self.spec.entity,
            "mode": self.mode,
            "dry_run": dry_run,
            "committed": committed,
            "mapping": self.mapping.mapping,
            "warnings": self.mapping.warnings,
            "rows": [row.report() for row in self.rows],
            "summary": self.summary(),
        }


def normalized_key(value: str) -> str:
    # Mirrors the API's case-insensitive uniqueness check: lower(trim(column)).
    return value.strip().lower()


def current_value(obj: Any, field_name: str) -> Any:
    return obj.metadata_ if field_name == "metadata" else getattr(obj, field_name)


def validation_errors(exc: ValidationError, provided: dict[str, Any]) -> list[dict[str, str]]:
    errors: list[dict[str, str]] = []
    for error in exc.errors():
        loc = error.get("loc") or ()
        field_name = str(loc[0]) if loc else "row"
        if error.get("type") == "missing":
            message = "is required"
        else:
            message = str(error.get("msg", "is invalid")).removeprefix("Value error, ")
            given = provided.get(field_name)
            if isinstance(given, str):
                message = f"{message} (got '{given}')"
        errors.append({"field": field_name, "message": message})
    return errors


def _existing_by_key(
    db: Session, spec: EntitySpec, keys: set[str], project_id: str | None
) -> dict[str, list[Any]]:
    if not keys:
        return {}
    column = getattr(spec.model, spec.key_field)
    found: dict[str, list[Any]] = {}
    key_list = sorted(keys)
    for start in range(0, len(key_list), 500):
        stmt = select(spec.model).where(
            func.lower(func.trim(column)).in_(key_list[start : start + 500])
        )
        if project_id is not None:
            stmt = stmt.where(spec.model.project_id == project_id)
        for obj in db.scalars(stmt):
            found.setdefault(normalized_key(getattr(obj, spec.key_field)), []).append(obj)
    return found


def _changes_for_create(validated: BaseModel) -> dict[str, list[Any]]:
    return {
        name: [None, value]
        for name, value in validated.model_dump(exclude={"project_id"}).items()
        if value not in (None, "", {})
    }


def plan_import(
    db: Session,
    spec: EntitySpec,
    table: Table,
    mode: ImportMode,
    *,
    project_id: str | None = None,
) -> ImportPlan:
    mapping = resolve_mapping(spec, table.headers)
    if spec.key_field not in {name for _, name in mapping.columns}:
        key = spec.field(spec.key_field)
        expected = ", ".join([key.name, *key.aliases])
        raise ImportFormatError(
            f"No column maps to {spec.key_field}; name one of: {expected}"
            if table.headers
            else "The input has no header row"
        )
    field_specs = {spec_field.name: spec_field for spec_field in spec.fields}

    parsed: list[tuple[int, dict[str, Any], list[dict[str, str]]]] = []
    for number, cells in table.rows:
        provided: dict[str, Any] = {}
        errors: list[dict[str, str]] = []
        for index, field_name in mapping.columns:
            text = cell_text(cells[index]) if index < len(cells) else None
            if text is None:
                continue  # Blank cell: keep the default on create, leave unchanged on update.
            try:
                provided[field_name] = coerce_cell(
                    field_specs[field_name], text, table.decimal_comma
                )
            except ValueError:
                errors.append({"field": field_name, "message": f"is not valid JSON (got '{text}')"})
        parsed.append((number, provided, errors))

    keys = {
        normalized_key(str(provided[spec.key_field]))
        for _, provided, _ in parsed
        if provided.get(spec.key_field)
    }
    existing = _existing_by_key(db, spec, keys, project_id)
    seen: dict[str, int] = {}
    plans: list[RowPlan] = []
    for number, provided, errors in parsed:
        raw_key = provided.get(spec.key_field)
        key = str(raw_key).strip() if raw_key is not None else None
        plan = RowPlan(row=number, action="error", key=key, errors=list(errors))
        plans.append(plan)
        if not key:
            plan.errors.append({"field": spec.key_field, "message": "is required"})
            continue
        norm = normalized_key(key)
        if norm in seen:
            plan.errors.append(
                {"field": spec.key_field, "message": f"duplicate of row {seen[norm]} in this file"}
            )
            continue
        seen[norm] = number
        if plan.errors:
            continue
        matches = existing.get(norm, [])
        if len(matches) > 1:
            plan.errors.append(
                {
                    "field": spec.key_field,
                    "message": f"matches {len(matches)} existing {spec.label} differing only "
                    "in case; rename them first",
                }
            )
            continue
        if not matches:
            _plan_create(plan, spec, provided, project_id)
            continue
        target = matches[0]
        plan.target = target
        if mode == "create_only":
            plan.errors.append(
                {
                    "field": spec.key_field,
                    "message": f"already exists as {getattr(target, spec.key_field)} "
                    "(use mode=upsert to update it)",
                }
            )
            continue
        _plan_update(plan, spec, provided, target)
    return ImportPlan(spec=spec, mode=mode, mapping=mapping, rows=plans, project_id=project_id)


def _plan_create(
    plan: RowPlan, spec: EntitySpec, provided: dict[str, Any], project_id: str | None
) -> None:
    data = dict(provided)
    if project_id is not None:
        data["project_id"] = project_id
    try:
        validated = spec.create_schema(**data)
    except ValidationError as exc:
        plan.errors.extend(validation_errors(exc, provided))
        return
    plan.action = "create"
    plan.create_data = validated.model_dump()
    plan.changes = _changes_for_create(validated)


def _plan_update(plan: RowPlan, spec: EntitySpec, provided: dict[str, Any], target: Any) -> None:
    # The key identifies the row; a case-only difference never renames the object.
    data = {name: value for name, value in provided.items() if name != spec.key_field}
    try:
        validated = spec.update_schema(**data)
    except ValidationError as exc:
        plan.errors.extend(validation_errors(exc, provided))
        return
    changes: dict[str, list[Any]] = {}
    for name, value in validated.model_dump(exclude_unset=True).items():
        old = current_value(target, name)
        if old != value:
            changes[name] = [old, value]
    plan.changes = changes
    plan.action = "update" if changes else "unchanged"


# --------------------------------------------------------------------------- bulk edit


def validate_bulk_changes(spec: EntitySpec, changes: dict[str, Any]) -> dict[str, Any]:
    """Validate a bulk edit through the update schema; returns the clean field values.

    Raises ValueError with a readable message for fields bulk edit may not touch.
    Pydantic ValidationError propagates for invalid values.
    """
    if not changes:
        raise ValueError("changes must name at least one field")
    disallowed = sorted(set(changes) - spec.bulk_fields)
    if disallowed:
        raise ValueError(
            f"Field(s) {', '.join(disallowed)} cannot be bulk edited; allowed: "
            + ", ".join(sorted(spec.bulk_fields))
        )
    cleared = sorted(
        name
        for name, value in changes.items()
        if value is None and name not in spec.nullable_fields
    )
    if cleared:
        raise ValueError(f"Field(s) {', '.join(cleared)} cannot be cleared")
    validated = spec.update_schema(**changes)
    return validated.model_dump(include=set(changes))


# --------------------------------------------------------------------------- templates


def template_rows(spec: EntitySpec) -> tuple[list[str], list[Any]]:
    fields = spec.template_fields
    return [spec_field.name for spec_field in fields], [spec_field.example for spec_field in fields]


def template_csv(spec: EntitySpec) -> bytes:
    headers, example = template_rows(spec)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerow(["" if value is None else value for value in example])
    # UTF-8 BOM so Excel opens the file as UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


def template_xlsx(spec: EntitySpec) -> bytes:
    headers, example = template_rows(spec)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.label.capitalize()
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="DDE3EA")
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = fill
        sheet.column_dimensions[cell.column_letter].width = max(12, len(header) + 4)
    sheet.append([None if value == "" else value for value in example])
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def table_to_csv(table: Table) -> str:
    """Re-emit a parsed table as comma-separated CSV text (for the line class importer)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(table.headers)
    for _, cells in table.rows:
        writer.writerow([cell_text(cell) or "" for cell in cells[: len(table.headers)]])
    return buffer.getvalue()
