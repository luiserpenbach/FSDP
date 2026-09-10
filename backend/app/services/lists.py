"""Engineering lists generated from the sheet index.

Instrument index, line list, valve list, equipment list, and tie-in list, per
drawing or per project, as JSON rows, CSV, or XLSX with a drawing header and a
zone column so every row can be found on the sheet.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Drawing, DrawingSheet, Part, Project, SheetItem, SheetLine
from app.services.sheet_index import drawing_index_rows

VALVE_CATEGORIES = {"valve", "regulator"}
INSTRUMENT_CATEGORIES = {"instrument"}
EQUIPMENT_CATEGORIES = {"equipment"}
CONNECTOR_CATEGORIES = {"connector"}


@dataclass(frozen=True)
class ListSpec:
    kind: str
    title: str
    source: str  # "items" or "lines"
    columns: tuple[tuple[str, str], ...]


LOCATION_COLUMNS: tuple[tuple[str, str], ...] = (("sheet_no", "Sheet"), ("zone", "Zone"))

LIST_SPECS: dict[str, ListSpec] = {
    "instrument": ListSpec(
        "instrument",
        "Instrument index",
        "items",
        (
            ("tag", "Tag"),
            ("symbol_name", "Type"),
            ("service", "Service"),
            ("line_number", "Line"),
            ("size", "Size"),
            ("mounting", "Mounting"),
            ("part_number", "Part"),
            ("dnp", "DNP"),
            ("notes", "Notes"),
        )
        + LOCATION_COLUMNS,
    ),
    "line": ListSpec(
        "line",
        "Line list",
        "lines",
        (
            ("line_number", "Line"),
            ("service", "Service"),
            ("line_type", "Type"),
            ("size", "Size"),
            ("spec", "Spec"),
            ("line_class", "Class"),
            ("from_tag", "From"),
            ("to_tag", "To"),
            ("design_pressure", "Design P"),
            ("design_temperature", "Design T"),
            ("operating_pressure", "Oper. P"),
            ("operating_temperature", "Oper. T"),
            ("insulation", "Insulation"),
            ("tracing", "Tracing"),
            ("length_m", "Length (m)"),
        )
        + LOCATION_COLUMNS,
    ),
    "valve": ListSpec(
        "valve",
        "Valve list",
        "items",
        (
            ("tag", "Tag"),
            ("symbol_name", "Type"),
            ("actuator", "Actuator"),
            ("size", "Size"),
            ("service", "Service"),
            ("line_number", "Line"),
            ("part_number", "Part"),
            ("dnp", "DNP"),
            ("notes", "Notes"),
        )
        + LOCATION_COLUMNS,
    ),
    "equipment": ListSpec(
        "equipment",
        "Equipment list",
        "items",
        (
            ("tag", "Tag"),
            ("name", "Name"),
            ("symbol_name", "Type"),
            ("nozzle_count", "Nozzles"),
            ("part_number", "Part"),
            ("notes", "Notes"),
        )
        + LOCATION_COLUMNS,
    ),
    "tie_in": ListSpec(
        "tie_in",
        "Tie-in list",
        "items",
        (
            ("tag", "Tag"),
            ("symbol_name", "Type"),
            ("ref", "Reference"),
            ("target", "Continues on"),
            ("service", "Service"),
            ("line_number", "Line"),
            ("size", "Size"),
            ("notes", "Notes"),
        )
        + LOCATION_COLUMNS,
    ),
}

LIST_KINDS = tuple(LIST_SPECS)


def _tag_sort_key(value: str | None) -> tuple:
    text = (value or "").upper()
    match = re.match(r"([A-Z]*)[\s\-_]*(\d*)(.*)", text)
    if not match:
        return (1, text)
    letters, digits, rest = match.groups()
    return (0 if letters or digits else 1, letters, int(digits) if digits else 0, rest)


def _item_matches(spec: ListSpec, item: SheetItem) -> bool:
    category = item.category or ""
    if spec.kind == "instrument":
        return category in INSTRUMENT_CATEGORIES
    if spec.kind == "valve":
        return category in VALVE_CATEGORIES
    if spec.kind == "equipment":
        return item.kind == "equipment" or category in EQUIPMENT_CATEGORIES
    if spec.kind == "tie_in":
        return category in CONNECTOR_CATEGORIES
    return False


def _location(sheet: DrawingSheet, drawing: Drawing) -> dict:
    return {
        "drawing_id": drawing.id,
        "drawing_number": drawing.number,
        "sheet_id": sheet.id,
        "sheet_no": sheet.sheet_no,
    }


def _item_row(item: SheetItem, sheet: DrawingSheet, drawing: Drawing, part: Part | None) -> dict:
    fields = dict(item.fields or {})
    return {
        **_location(sheet, drawing),
        "item_id": item.item_id,
        "kind": item.kind,
        "category": item.category,
        "tag": item.tag or item.label,
        "name": item.label or item.symbol_name,
        "symbol_key": item.symbol_key,
        "symbol_name": item.symbol_name,
        "zone": item.zone,
        "part_id": item.part_id,
        "part_number": part.part_number if part else None,
        "dnp": "DNP" if item.dnp else "",
        "spare": item.spare,
        "service": fields.get("service"),
        "line_number": fields.get("line_number"),
        "size": fields.get("size"),
        "actuator": fields.get("actuator"),
        "mounting": fields.get("mounting"),
        "nozzle_count": fields.get("nozzle_count"),
        "ref": fields.get("ref"),
        "target": fields.get("target"),
        "notes": fields.get("notes"),
    }


def _line_row(line: SheetLine, sheet: DrawingSheet, drawing: Drawing) -> dict:
    return {
        **_location(sheet, drawing),
        "item_id": line.line_id,
        "kind": "line",
        "line_number": line.line_number,
        "service": line.service,
        "line_type": line.line_type,
        "size": line.size,
        "spec": line.spec,
        "line_class": line.line_class,
        "from_item": line.from_item,
        "from_tag": line.from_tag,
        "to_item": line.to_item,
        "to_tag": line.to_tag,
        "zone": line.zone,
        "length_mm": line.length_mm,
        "length_m": line.length_m,
        "connection_count": line.connection_count,
        "tee_count": line.tee_count,
        "design_pressure": line.design_pressure,
        "design_temperature": line.design_temperature,
        "operating_pressure": line.operating_pressure,
        "operating_temperature": line.operating_temperature,
        "insulation": line.insulation,
        "tracing": line.tracing,
    }


def list_rows(db: Session, kind: str, drawings: list[Drawing]) -> list[dict]:
    """Rows of the list `kind` over the given drawings, sorted by tag / line number."""
    spec = LIST_SPECS[kind]
    items, lines = drawing_index_rows(db, drawings)
    if spec.source == "lines":
        rows = [_line_row(line, sheet, drawing) for line, sheet, drawing in lines]
        rows.sort(key=lambda row: (row["drawing_number"], _tag_sort_key(row["line_number"])))
        return rows
    part_ids = {item.part_id for item, _, _ in items if item.part_id}
    parts = (
        {part.id: part for part in db.query(Part).filter(Part.id.in_(part_ids))} if part_ids else {}
    )
    rows = [
        _item_row(item, sheet, drawing, parts.get(item.part_id) if item.part_id else None)
        for item, sheet, drawing in items
        if _item_matches(spec, item)
    ]
    rows.sort(key=lambda row: (row["drawing_number"], _tag_sort_key(row["tag"])))
    return rows


def list_columns(kind: str, scope: str) -> list[tuple[str, str]]:
    columns = list(LIST_SPECS[kind].columns)
    if scope == "project":
        columns.insert(0, ("drawing_number", "Drawing"))
    return columns


def list_header(kind: str, scope: str, drawing: Drawing | None, project: Project) -> dict:
    header: dict = {
        "list": LIST_SPECS[kind].title,
        "project": project.name,
        "generated": datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC"),
    }
    if drawing is not None:
        current = drawing.revisions[-1] if drawing.revisions else None
        header.update(
            {
                "drawing_number": drawing.number,
                "title": drawing.title.replace("\n", " "),
                "revision": current.label if current else "-",
                "status": drawing.status,
                "sheets": len(drawing.sheets),
            }
        )
    return header


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def _csv_safe(value: str) -> str:
    # Guard spreadsheet formula injection when the CSV is opened in Excel.
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def rows_to_csv(header: dict, columns: list[tuple[str, str]], rows: list[dict]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for key, value in header.items():
        writer.writerow([key.replace("_", " ").title(), _csv_safe(_cell(value))])
    writer.writerow([])
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([_csv_safe(_cell(row.get(key))) for key, _ in columns])
    return buffer.getvalue()


def rows_to_xlsx(header: dict, columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (header.get("list") or "List")[:31]
    bold = Font(bold=True)
    for key, value in header.items():
        sheet.append([key.replace("_", " ").title(), _cell(value)])
        sheet.cell(row=sheet.max_row, column=1).font = bold
    sheet.append([])
    sheet.append([label for _, label in columns])
    header_row = sheet.max_row
    fill = PatternFill("solid", fgColor="DDE3EA")
    for column_index in range(1, len(columns) + 1):
        cell = sheet.cell(row=header_row, column=column_index)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(
            [
                row.get(key) if isinstance(row.get(key), int | float) else _cell(row.get(key))
                for key, _ in columns
            ]
        )
    for column_index, (key, label) in enumerate(columns, start=1):
        width = max([len(label)] + [len(_cell(row.get(key))) for row in rows] + [4])
        sheet.column_dimensions[get_column_letter(column_index)].width = min(48, width + 2)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def list_filename(kind: str, scope_name: str, fmt: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", scope_name).strip("-.") or "list"
    return f"{slug}-{kind.replace('_', '-')}-list.{fmt}"
