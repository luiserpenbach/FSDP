"""Requirements import and export (CSV/XLSX) and the coverage report."""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Hazard,
    Project,
    Requirement,
    RequirementEvidence,
    SheetItem,
    TraceLink,
)
from app.schemas import REQUIREMENT_CATEGORIES, clean_verification_method
from app.services.hazards import resolve_controls

IMPORT_FIELDS = (
    "key",
    "title",
    "text",
    "category",
    "verification_method",
    "status",
    "owner",
    "parent_key",
    "rationale",
    "source_ref",
)

# Header spellings recognised without an explicit mapping.
DEFAULT_ALIASES = {
    "key": {"key", "id", "req id", "requirement id", "identifier"},
    "title": {"title", "name", "summary"},
    "text": {"text", "requirement", "requirement text", "statement", "shall"},
    "category": {"category", "type", "requirement type", "requirement_type"},
    "verification_method": {
        "verification",
        "verification method",
        "verification_method",
        "method",
    },
    "status": {"status", "state"},
    "owner": {"owner", "responsible"},
    "parent_key": {
        "parent",
        "parent key",
        "parent_key",
        "derived from",
        "derives from",
        "source req",
    },
    "rationale": {"rationale", "justification"},
    "source_ref": {"source", "source ref", "source_ref", "reference", "clause"},
}

EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("key", "Key"),
    ("title", "Title"),
    ("text", "Text"),
    ("category", "Category"),
    ("verification_method", "Verification method"),
    ("verification_status", "Verification status"),
    ("status", "Status"),
    ("owner", "Owner"),
    ("parent_key", "Derives from"),
    ("safety_critical", "Safety critical"),
    ("rationale", "Rationale"),
    ("source_ref", "Source"),
    ("constraint", "Design rule"),
    ("revision", "Revision"),
]

MATRIX_COLUMNS: list[tuple[str, str]] = [
    ("key", "Key"),
    ("title", "Title"),
    ("category", "Category"),
    ("verification_method", "Method"),
    ("owner", "Owner"),
    ("verification_status", "Verification status"),
    ("verdict", "DRC verdict"),
    ("checked", "Items checked"),
    ("failed", "Items failed"),
    ("evidence", "Evidence"),
    ("hazards", "Hazards"),
    ("drawings", "Drawings"),
    ("safety_critical", "Safety critical"),
]


def _table_from_rows(raw_rows: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    """Header and rows from a grid. Our own exports start with a key/value block
    and a blank line before the column header; that block is skipped."""
    rows = [[cell.strip() for cell in row] for row in raw_rows]
    start = 0
    for index, row in enumerate(rows[:12]):
        if not any(row) and index + 1 < len(rows):
            start = index + 1
            break
    body = rows[start:]
    if not body:
        return [], []
    header = body[0]
    records = []
    for cells in body[1:]:
        if not any(cells):
            continue
        records.append({header[i]: cells[i] for i in range(min(len(header), len(cells)))})
    return [name for name in header], records


def parse_table(filename: str, payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """Header names and rows (as strings) from a CSV or XLSX upload."""
    lowered = (filename or "").lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        grid = [
            ["" if value is None else str(value) for value in values]
            for values in sheet.iter_rows(values_only=True)
        ]
        return _table_from_rows(grid)
    text = payload.decode("utf-8-sig", errors="replace")
    grid = [list(row) for row in csv.reader(io.StringIO(text))]
    return _table_from_rows(grid)


def resolve_mapping(header: list[str], mapping: dict[str, str] | None) -> dict[str, str]:
    """Column -> field. Explicit mapping first, then header aliases."""
    resolved: dict[str, str] = {}
    for column, field in (mapping or {}).items():
        if field in IMPORT_FIELDS and column in header:
            resolved[column] = field
    taken = set(resolved.values())
    for column in header:
        if column in resolved:
            continue
        lowered = column.strip().lower()
        for field, aliases in DEFAULT_ALIASES.items():
            if field not in taken and lowered in aliases:
                resolved[column] = field
                taken.add(field)
                break
    return resolved


def import_requirements(
    db: Session,
    project: Project,
    header: list[str],
    rows: list[dict[str, str]],
    mapping: dict[str, str] | None,
    *,
    dry_run: bool,
    update_existing: bool,
) -> dict[str, Any]:
    columns = resolve_mapping(header, mapping)
    fields_present = set(columns.values())
    errors: list[str] = []
    if "key" not in fields_present:
        errors.append("No column maps to 'key'")
    if "title" not in fields_present and "text" not in fields_present:
        errors.append("No column maps to 'title' or 'text'")
    if errors:
        return {
            "dry_run": dry_run,
            "mapping": columns,
            "created": 0,
            "updated": 0,
            "skipped": 0,
            "errors": errors,
            "preview": [],
        }

    existing = {
        requirement.key: requirement
        for requirement in db.scalars(
            select(Requirement).where(Requirement.project_id == project.id)
        )
    }
    seen: set[str] = set()
    planned: list[dict[str, Any]] = []
    for index, raw in enumerate(rows, start=2):
        record = {field: (raw.get(column) or "").strip() for column, field in columns.items()}
        key = record.get("key", "")
        if not key:
            errors.append(f"row {index}: missing key")
            continue
        if key in seen:
            errors.append(f"row {index}: duplicate key {key} in the file")
            continue
        seen.add(key)
        title = record.get("title") or record.get("text", "")[:200]
        text = record.get("text") or title
        category = (record.get("category") or "functional").strip().lower()
        if category not in REQUIREMENT_CATEGORIES:
            category = "functional"
        entry = {
            "row": index,
            "key": key,
            "title": title,
            "text": text,
            "category": category,
            "verification_method": clean_verification_method(record.get("verification_method"))
            or None,
            "status": (record.get("status") or "draft").strip().lower() or "draft",
            "owner": record.get("owner") or None,
            "parent_key": record.get("parent_key") or None,
            "rationale": record.get("rationale") or None,
            "source_ref": record.get("source_ref") or None,
            "action": "update" if key in existing else "create",
        }
        if key in existing and not update_existing:
            entry["action"] = "skip"
        planned.append(entry)

    created = updated = skipped = 0
    by_key: dict[str, Requirement] = dict(existing)
    if not dry_run:
        for entry in planned:
            if entry["action"] == "skip":
                skipped += 1
                continue
            values = {
                "title": entry["title"],
                "text": entry["text"],
                "requirement_type": entry["category"],
                "category": entry["category"],
                "verification_method": entry["verification_method"],
                "status": entry["status"],
                "owner": entry["owner"],
                "rationale": entry["rationale"],
                "source_ref": entry["source_ref"],
            }
            if entry["action"] == "update":
                requirement = by_key[entry["key"]]
                for field, value in values.items():
                    if value is not None:
                        setattr(requirement, field, value)
                updated += 1
            else:
                requirement = Requirement(project_id=project.id, key=entry["key"], **values)
                db.add(requirement)
                by_key[entry["key"]] = requirement
                created += 1
        db.flush()
        for entry in planned:
            parent_key = entry["parent_key"]
            if entry["action"] == "skip" or not parent_key:
                continue
            parent = by_key.get(parent_key)
            if parent is None:
                errors.append(f"row {entry['row']}: parent {parent_key} not found")
                continue
            requirement = by_key[entry["key"]]
            if parent.id != requirement.id:
                requirement.parent_id = parent.id
        db.flush()
    else:
        for entry in planned:
            if entry["action"] == "skip":
                skipped += 1
            elif entry["action"] == "update":
                updated += 1
            else:
                created += 1
            parent_key = entry["parent_key"]
            if parent_key and parent_key not in existing and parent_key not in seen:
                errors.append(f"row {entry['row']}: parent {parent_key} not found")
    return {
        "dry_run": dry_run,
        "mapping": columns,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "preview": planned[:20],
    }


def export_rows(db: Session, project: Project) -> list[dict[str, Any]]:
    requirements = list(
        db.scalars(
            select(Requirement)
            .where(Requirement.project_id == project.id)
            .order_by(Requirement.key)
        )
    )
    keys = {requirement.id: requirement.key for requirement in requirements}
    rows = []
    for requirement in requirements:
        constraint = requirement.constraint or {}
        rows.append(
            {
                "key": requirement.key,
                "title": requirement.title,
                "text": requirement.text,
                "category": requirement.category,
                "verification_method": requirement.verification_method,
                "verification_status": requirement.verification_status,
                "status": requirement.status,
                "owner": requirement.owner,
                "parent_key": keys.get(requirement.parent_id or ""),
                "safety_critical": "yes" if requirement.safety_critical else "no",
                "rationale": requirement.rationale,
                "source_ref": requirement.source_ref,
                "constraint": (
                    f"{constraint.get('kind')} {' '.join(constraint.get('values') or [])}".strip()
                    if constraint
                    else ""
                ),
                "revision": requirement.revision,
            }
        )
    return rows


def matrix_rows(matrix: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in matrix["rows"]:
        rows.append(
            {
                **{key: row.get(key) for key, _ in MATRIX_COLUMNS},
                "evidence": ", ".join(
                    f"{count} {kind}" for kind, count in sorted((row.get("evidence") or {}).items())
                ),
                "hazards": ", ".join(row.get("hazards") or []),
                "drawings": "; ".join(
                    f"{entry['drawing_number']} "
                    f"(sheets {', '.join(str(s) for s in entry['sheets'])})"
                    for entry in row.get("drawings") or []
                ),
                "safety_critical": "yes" if row.get("safety_critical") else "no",
            }
        )
    return rows


def coverage(db: Session, project: Project) -> dict[str, Any]:
    requirements = list(
        db.scalars(
            select(Requirement)
            .where(Requirement.project_id == project.id)
            .order_by(Requirement.key)
        )
    )
    requirement_ids = [requirement.id for requirement in requirements]
    traced: set[str] = set()
    if requirement_ids:
        for (source_id,) in db.execute(
            select(TraceLink.source_id).where(
                TraceLink.source_type == "requirement",
                TraceLink.source_id.in_(requirement_ids),
                TraceLink.target_type.in_(("drawing", "sheet_item", "sheet_line", "component")),
            )
        ):
            traced.add(source_id)
    with_evidence: set[str] = set()
    if requirement_ids:
        for (requirement_id,) in db.execute(
            select(RequirementEvidence.requirement_id).where(
                RequirementEvidence.requirement_id.in_(requirement_ids)
            )
        ):
            with_evidence.add(requirement_id)

    def brief(requirement: Requirement) -> dict[str, str]:
        return {"id": requirement.id, "key": requirement.key, "title": requirement.title}

    hazards = list(
        db.scalars(select(Hazard).where(Hazard.project_id == project.id).order_by(Hazard.key))
    )
    hazards_without_controls = []
    uncovered_controls = []
    for hazard in hazards:
        controls = resolve_controls(db, hazard)
        if not controls and hazard.status not in {"closed"}:
            hazards_without_controls.append(
                {"id": hazard.id, "key": hazard.key, "title": hazard.title}
            )
        for control in controls:
            if control["type"] == "sheet_item" and not control["covered"]:
                item = db.get(SheetItem, control["id"])
                uncovered_controls.append(
                    {
                        "hazard_id": hazard.id,
                        "hazard_key": hazard.key,
                        "link_id": control["link_id"],
                        "item_id": control["id"],
                        "tag": control["label"],
                        "sheet_id": item.sheet_id if item else None,
                    }
                )
    return {
        "project_id": project.id,
        "totals": {
            "requirements": len(requirements),
            "hazards": len(hazards),
            "traced": len(traced),
            "with_evidence": len(with_evidence),
        },
        "untraced_requirements": [brief(r) for r in requirements if r.id not in traced],
        "critical_without_evidence": [
            brief(r) for r in requirements if r.safety_critical and r.id not in with_evidence
        ],
        "hazards_without_controls": hazards_without_controls,
        "uncovered_hardware_controls": uncovered_controls,
    }
